"""
Herramienta de Contabilidad
"""
import os
from core.tool_registry import register
from database.aronium import registrar_gasto, get_gastos, get_balance_mes, get_valor_inventario, get_margen_productos
from services.ollama_service import analizar_ollama

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def gasto_tool(data):
    if isinstance(data, str): return "Error: datos invalidos"
    descripcion = data.get("descripcion", "")
    monto = data.get("monto", 0)
    categoria = data.get("categoria", "general")
    registrar_gasto(descripcion, monto, categoria)
    return f"Gasto registrado: {descripcion} - Q{monto}"

register("contabilidad.gasto", gasto_tool)

def gastos_tool(data):
    gastos = get_gastos()
    if not gastos:
        return "No hay gastos registrados este mes."
    total = sum(float(g["Monto"]) for g in gastos)
    resp = f"Gastos del mes (Q{total:.2f} total):\n"
    for g in gastos:
        resp += f"- {g['Descripcion']}: Q{g['Monto']} ({g['Fecha'][:10]})\n"
    return resp

register("contabilidad.gastos", gastos_tool)

def balance_tool(data):
    b = get_balance_mes()
    resp = f"Balance del mes:\n"
    resp += f"Ingresos por ventas: Q{b['ingresos']}\n"
    resp += f"Costo de productos: Q{b['costos_productos']}\n"
    resp += f"Gastos operativos: Q{b['gastos_operativos']}\n"
    resp += f"Comision POS (6%): Q{b['comision_pos']}\n"
    resp += f"Utilidad bruta: Q{b['utilidad_bruta']}\n"
    resp += f"Utilidad neta: Q{b['utilidad_neta']}\n"
    resp += f"---\n"
    resp += f"IVA estimado (12%): Q{b['iva_estimado']}\n"
    resp += f"ISR estimado (5%): Q{b['isr_estimado']}"
    return resp

register("contabilidad.balance", balance_tool)

def inventario_valor_tool(data):
    inv = get_valor_inventario()
    resp = f"Valor total del inventario: Q{inv['total_inventario']}\n\n"
    for p in inv['productos']:
        resp += f"- {p['Name']}: {int(p['Stock'])} uds x Q{p['Cost']} = Q{p['ValorTotal']:.2f}\n"
    return resp

register("contabilidad.inventario_valor", inventario_valor_tool)

def margen_tool(data):
    productos = get_margen_productos()
    if not productos:
        return "No hay productos con precio y costo registrados."
    resp = "Margen de ganancia por producto:\n"
    for p in productos:
        resp += f"- {p['Name']}: Q{p['Cost']} costo → Q{p['Price']} precio = {p['Margen']}% margen\n"
    return resp

register("contabilidad.margen", margen_tool)


def sugerencias_precios_tool(data):
    productos = get_margen_productos()
    if not productos:
        return "No hay productos con precio y costo registrados para analizar."

    balance = get_balance_mes()

    COMISION_POS = 0.06
    MARGEN_OBJETIVO = 0.30  # margen minimo saludable, sobre el precio de venta

    analizados = []
    for p in productos:
        precio = float(p["Price"])
        costo = float(p["Cost"])
        if precio <= 0:
            continue
        # Margen real = ganancia sobre el precio de venta (no sobre el costo).
        # Es distinto al campo "Margen" de get_margen_productos(), que es markup sobre costo.
        margen_real = round((precio - costo) / precio * 100, 1)
        bien = margen_real >= MARGEN_OBJETIVO * 100
        precio_sugerido = None
        if not bien:
            # Precio minimo para que, ya pagado el costo y la comision de la POS (6%),
            # quede al menos un 30% de margen sobre ese precio.
            divisor = 1 - COMISION_POS - MARGEN_OBJETIVO
            if divisor > 0:
                precio_sugerido = round(costo / divisor, 2)
        analizados.append({
            "nombre": p["Name"].strip(),
            "costo": costo,
            "precio_actual": precio,
            "margen_real": margen_real,
            "bien": bien,
            "precio_sugerido": precio_sugerido
        })

    if not analizados:
        return "No hay productos con precio valido para analizar margenes."

    buenos = [a for a in analizados if a["bien"]]
    bajos = [a for a in analizados if not a["bien"]]

    def _resumen_sin_ia():
        resp = f"Analisis de precios ({len(analizados)} productos, margen objetivo 30%):\n\n"
        if buenos:
            resp += f"BIEN ({len(buenos)}):\n"
            for a in buenos:
                resp += f"- {a['nombre']}: {a['margen_real']}% margen (Q{a['precio_actual']:.2f})\n"
        if bajos:
            resp += f"\nNECESITAN AJUSTE ({len(bajos)}):\n"
            for a in bajos:
                sugerido = f"Q{a['precio_sugerido']:.2f}" if a["precio_sugerido"] else "N/A"
                resp += f"- {a['nombre']}: {a['margen_real']}% margen actual (Q{a['precio_actual']:.2f}) -> sugerido {sugerido}\n"
        return resp

    tabla_buenos = "\n".join(
        f"- {a['nombre']}: costo Q{a['costo']:.2f}, precio Q{a['precio_actual']:.2f}, margen real {a['margen_real']}%"
        for a in buenos
    ) or "(ninguno)"

    tabla_bajos = "\n".join(
        f"- {a['nombre']}: costo Q{a['costo']:.2f}, precio actual Q{a['precio_actual']:.2f}, "
        f"margen real {a['margen_real']}%, precio sugerido Q{a['precio_sugerido']:.2f}"
        for a in bajos
    ) or "(ninguno)"

    prompt = f"""Eres el analista de precios de Mi Cocina Electrodomesticos, una tienda de electrodomesticos y repuestos en Guatemala.

El margen real de cada producto se calcula como (Precio - Costo) / Precio * 100.
El objetivo del negocio es un margen minimo de 30% sobre el precio de venta para ser sostenible.
Ya se clasifico cada producto en dos grupos usando ese 30% como corte exacto. NO reclasifiques
ningun producto ni le apliques un margen de seguridad adicional: si esta en el grupo "SALUDABLE"
es porque ya cumple el objetivo (aunque este apenas por encima del 30%), y si esta en "NECESITA AJUSTE"
el precio sugerido ya fue calculado matematicamente (cubre costo + comision de la POS 6% + 30% de margen).
No inventes ni recalcules ningun precio sugerido para los productos del grupo SALUDABLE.

Factores del negocio para contextualizar tu analisis (no para recalcular los precios):
- Comision de la POS (procesador de pagos): 6% de cada venta
- IVA: 12% (normalmente ya incluido en el precio de venta al publico en Guatemala)
- Gastos operativos de este mes: Q{balance['gastos_operativos']:.2f}
- Ingresos del mes: Q{balance['ingresos']:.2f}

GRUPO SALUDABLE (margen >= 30%, no necesitan cambio de precio):
{tabla_buenos}

GRUPO NECESITA AJUSTE (margen < 30%, precio sugerido ya calculado):
{tabla_bajos}

Responde en espanol, de forma clara y directa, con:
1. Un resumen: cuantos productos estan con margen saludable y cuantos necesitan ajuste
2. Cuales productos estan bien (mencionalos brevemente, sin extenderte, y sin sugerirles cambios)
3. Cuales necesitan ajuste de precio y por cuanto deben subir (usa exactamente los precios sugeridos que te di, no otros)
4. Una recomendacion practica para una tienda pequena en Guatemala: si el ajuste es grande, sugiere subir el precio de forma gradual para no espantar clientes. Si el grupo "NECESITA AJUSTE" esta vacio, dilo claramente y no inventes ajustes."""

    texto = analizar_ollama(prompt)
    return texto if texto else _resumen_sin_ia()

register("contabilidad.sugerencias_precios", sugerencias_precios_tool)


def equilibrio_tool(data):
    from database.aronium import get_gastos, get_pagos_fijos, get_margen_productos, get_balance_mes
    from datetime import datetime
    
    # Gastos operativos del mes
    gastos = get_gastos()
    total_gastos_variables = sum(float(g["Monto"]) for g in gastos)
    
    # Pagos fijos mensuales
    pagos = get_pagos_fijos()
    total_pagos_fijos = sum(float(p["Monto"] or 0) for p in pagos if p["Frecuencia"] == "mensual")

    total_costos_fijos = total_gastos_variables + total_pagos_fijos
    
    # Margen promedio de productos
    productos = get_margen_productos()
    if not productos:
        return "No hay productos con precio y costo para calcular."
    
    margenes = [float(p["Margen"]) for p in productos if float(p["Margen"]) > 0]
    margen_promedio = sum(margenes) / len(margenes) if margenes else 0
    
    # Precio promedio de venta
    precio_promedio = sum(float(p["Price"]) for p in productos if float(p["Price"]) > 0) / len([p for p in productos if float(p["Price"]) > 0])
    costo_promedio = sum(float(p["Cost"]) for p in productos if float(p["Cost"]) > 0) / len([p for p in productos if float(p["Cost"]) > 0])
    
    margen_unitario = precio_promedio - costo_promedio
    comision_pos_unitaria = precio_promedio * 0.06
    contribucion_unitaria = margen_unitario - comision_pos_unitaria
    
    if contribucion_unitaria <= 0:
        return "No se puede calcular el punto de equilibrio con los margenes actuales."
    
    unidades_equilibrio = total_costos_fijos / contribucion_unitaria
    ventas_equilibrio = unidades_equilibrio * precio_promedio
    
    resp = "Punto de equilibrio del mes:\n"
    resp += f"Gastos operativos registrados: Q{total_gastos_variables:.2f}\n"
    resp += f"Pagos fijos mensuales: Q{total_pagos_fijos:.2f}\n"
    resp += f"Precio promedio de venta: Q{precio_promedio:.2f}\n"
    resp += f"Costo promedio de producto: Q{costo_promedio:.2f}\n"
    resp += f"Contribucion por unidad: Q{contribucion_unitaria:.2f}\n"
    resp += f"---\n"
    resp += f"Necesitas vender {unidades_equilibrio:.0f} unidades para cubrir gastos\n"
    resp += f"O generar Q{ventas_equilibrio:.2f} en ventas\n"
    ventas_mes_actual = get_balance_mes()["ingresos"]
    resp += f"Llevas Q{ventas_mes_actual:.2f} vendidos este mes"
    return resp

register("contabilidad.equilibrio", equilibrio_tool)


def analisis_completo_tool(data):
    from database.aronium import get_datos_completos_empresa

    datos = get_datos_completos_empresa()
    
    prompt = f"""Eres el analista financiero y estratega de negocios de Mi Cocina Electrodomesticos Guatemala.
Tienda pequena de electrodomesticos y repuestos. Analiza los datos reales del negocio y da recomendaciones especificas.

DATOS DEL MES ({datos['fecha']}):

VENTAS:
- Total vendido: Q{datos['ventas']['total_mes']}
- Transacciones: {datos['ventas']['transacciones']}
- Comision POS: Q{datos['ventas']['comision_pos']}
- Ganancia estimada: Q{datos['ventas']['ganancia_estimada']}

BALANCE:
- Ingresos: Q{datos['balance']['ingresos']}
- Costos productos: Q{datos['balance']['costos_productos']}
- Gastos operativos: Q{datos['balance']['gastos_operativos']}
- Utilidad neta: Q{datos['balance']['utilidad_neta']}
- IVA estimado: Q{datos['balance']['iva_estimado']}
- ISR estimado: Q{datos['balance']['isr_estimado']}

INVENTARIO:
- Valor total invertido: Q{datos['inventario']['valor_total']}
- Productos activos: {datos['inventario']['productos']}
- Productos con stock bajo: {datos['alertas_stock']}

MARGENES DE PRODUCTOS:
{chr(10).join([f"- {p['Name']}: {p['Margen']}% margen (Q{p['Cost']} costo -> Q{p['Price']} precio)" for p in datos['margenes']])}

CLIENTES:
- Total clientes: {datos['clientes']['total']}
- Clientes inactivos: {datos['clientes']['inactivos']}
- Mejores clientes: {', '.join([c['Nombre'] for c in datos['clientes']['mejores']]) if datos['clientes']['mejores'] else 'Sin datos'}

Con base en estos datos REALES proporciona:

1. DIAGNOSTICO: Estado actual del negocio (3-4 oraciones)
2. FORTALEZAS: Que esta funcionando bien
3. ALERTAS: Que necesita atencion inmediata
4. META DEL PROXIMO MES: Cuanto deberia vender para crecer un 15-20%
5. ESTRATEGIAS CONCRETAS: 3 acciones especificas para lograr la meta (considera redes sociales, precios, clientes inactivos, stock)
6. PROYECCION: Si implementa las estrategias, que resultado puede esperar

Responde en espanol, de forma directa y practica para un negocio guatemalteco pequeno."""

    texto = analizar_ollama(prompt)
    return texto if texto else "No pude generar el analisis: Ollama no esta disponible ahora mismo."

register("contabilidad.analisis_completo", analisis_completo_tool)


def excel_contabilidad_tool(data):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from datetime import datetime
    from database.aronium import get_balance_mes, get_gastos, get_margen_productos, get_valor_inventario, get_reporte_periodo, nimbo_connect
    
    now = datetime.now()
    mes = now.month
    anio = now.year
    
    wb = openpyxl.Workbook()
    
    # ── HOJA 1: BALANCE ──
    ws1 = wb.active
    ws1.title = "Balance Mensual"
    
    b = get_balance_mes()
    ws1.append(["Mi Cocina Electrodomesticos Guatemala"])
    ws1.append([f"Balance Mensual - {now.strftime('%B %Y')}"])
    ws1.append([])
    ws1.append(["Concepto", "Monto Q"])
    ws1.append(["INGRESOS", ""])
    ws1.append(["Ventas totales", b["ingresos"]])
    ws1.append(["", ""])
    ws1.append(["COSTOS Y GASTOS", ""])
    ws1.append(["Costo de productos vendidos", b["costos_productos"]])
    ws1.append(["Gastos operativos", b["gastos_operativos"]])
    ws1.append(["Comision POS (6%)", b["comision_pos"]])
    ws1.append(["", ""])
    ws1.append(["UTILIDAD BRUTA", b["utilidad_bruta"]])
    ws1.append(["UTILIDAD NETA", b["utilidad_neta"]])
    ws1.append(["", ""])
    ws1.append(["IMPUESTOS ESTIMADOS", ""])
    ws1.append(["IVA (12%)", b["iva_estimado"]])
    ws1.append(["ISR (5%)", b["isr_estimado"]])
    
    # ── HOJA 2: GASTOS ──
    ws2 = wb.create_sheet("Gastos")
    ws2.append(["Fecha", "Descripcion", "Categoria", "Monto Q"])
    gastos = get_gastos()
    for g in gastos:
        ws2.append([g["Fecha"][:10], g["Descripcion"], g["Categoria"], float(g["Monto"])])
    ws2.append(["", "", "TOTAL", sum(float(g["Monto"]) for g in gastos)])
    
    # ── HOJA 3: VENTAS ──
    ws3 = wb.create_sheet("Ventas")
    ws3.append(["Fecha", "Descripcion", "Total Q", "Cliente"])
    conn = nimbo_connect()
    cur = conn.cursor()
    cur.execute("""SELECT v.Fecha, v.Descripcion, v.Total, IFNULL(c.Nombre,'Sin cliente')
        FROM Venta v LEFT JOIN Cliente c ON v.ClienteId=c.Id
        WHERE v.Mes=? AND v.Anio=?""", (mes, anio))
    for r in cur.fetchall():
        ws3.append(list(r))
    conn.close()
    
    # ── HOJA 4: INVENTARIO ──
    ws4 = wb.create_sheet("Inventario")
    ws4.append(["Producto", "Stock", "Costo Unit Q", "Precio Venta Q", "Margen %", "Valor Invertido Q"])
    inv = get_valor_inventario()
    margenes = {p["Name"].strip(): p for p in get_margen_productos()}
    for p in inv["productos"]:
        nombre = p["Name"].strip()
        m = margenes.get(nombre, {})
        ws4.append([
            nombre,
            int(p["Stock"]),
            float(p["Cost"]),
            float(m.get("Price", 0)),
            float(m.get("Margen", 0)),
            float(p["ValorTotal"])
        ])
    ws4.append(["TOTAL", "", "", "", "", inv["total_inventario"]])
    
    filename = os.path.join(BASE_DIR, f"contabilidad_{now.strftime('%Y_%m')}.xlsx")
    wb.save(filename)
    return f"DOWNLOAD:contabilidad_{now.strftime('%Y_%m')}.xlsx"

register("contabilidad.excel", excel_contabilidad_tool)


def comparacion_meses_tool(data):
    from database.aronium import get_comparacion_meses
    meses = get_comparacion_meses(6)
    
    resp = "Comparacion ultimos 6 meses:\n\n"
    resp += f"{'Mes':<12} {'Ventas':>8} {'Total':>10} {'Gastos':>10} {'Utilidad':>10}\n"
    resp += "-" * 52 + "\n"
    
    for m in meses:
        resp += f"{m['nombre']:<12} {m['ventas']:>8} Q{m['total']:>8.2f} Q{m['gastos']:>8.2f} Q{m['utilidad']:>8.2f}\n"
    
    # Calcular crecimiento
    if len(meses) >= 2 and meses[-2]["total"] > 0:
        crecimiento = ((meses[-1]["total"] - meses[-2]["total"]) / meses[-2]["total"]) * 100
        if crecimiento > 0:
            resp += f"\nCrecimiento vs mes anterior: +{crecimiento:.1f}%"
        else:
            resp += f"\nCambio vs mes anterior: {crecimiento:.1f}%"
    
    return resp

register("contabilidad.comparacion", comparacion_meses_tool)


def flujo_proyectado_tool(data):
    from database.aronium import get_flujo_caja_proyectado
    f = get_flujo_caja_proyectado()
    signo = "sobraran" if f["positivo"] else "faltaran"
    resp = f"Proyeccion de flujo de caja para {f['mes_proyectado']}:\n\n"
    resp += f"Ingresos proyectados: Q{f['ingresos']['total']:.2f}\n"
    resp += f"Egresos proyectados: Q{f['egresos']['total']:.2f}\n"
    resp += "---\n"
    if f["positivo"]:
        resp += f"Flujo neto proyectado: +Q{f['flujo_neto']:.2f} ({signo} Q{abs(f['flujo_neto']):.2f})"
    else:
        resp += f"Flujo neto proyectado: -Q{abs(f['flujo_neto']):.2f}\n"
        resp += f"ALERTA: te {signo} Q{abs(f['flujo_neto']):.2f} para cubrir los egresos del mes."
    return resp

register("contabilidad.flujo_proyectado", flujo_proyectado_tool)


def flujo_proyectado_detalle_tool(data):
    from database.aronium import get_flujo_caja_proyectado
    f = get_flujo_caja_proyectado()
    ing = f["ingresos"]
    egr = f["egresos"]

    resp = f"Proyeccion financiera detallada para {f['mes_proyectado']}:\n\n"

    resp += "INGRESOS PROYECTADOS:\n"
    resp += f"- Promedio de ventas (ultimos 3 meses): Q{ing['promedio_ventas_3_meses']:.2f}\n"
    tendencia_txt = "subiendo" if ing["tendencia_pct"] > 0 else ("bajando" if ing["tendencia_pct"] < 0 else "estable")
    resp += f"- Tendencia: {tendencia_txt} ({ing['tendencia_pct']:+.1f}% promedio mensual, ajuste Q{ing['ajuste_tendencia']:+.2f})\n"
    resp += f"- Ventas proyectadas: Q{ing['ventas_proyectadas']:.2f}\n"
    resp += f"- Cuentas por cobrar que vencen el proximo mes: Q{ing['cuentas_por_cobrar']:.2f} ({ing['cantidad_cuentas_cobrar']} cuentas)\n"
    resp += f"Total ingresos: Q{ing['total']:.2f}\n\n"

    resp += "EGRESOS PROYECTADOS:\n"
    resp += f"- Pagos fijos ({egr['cantidad_pagos_fijos']}): Q{egr['pagos_fijos']:.2f}\n"
    if egr["pagos_fijos_sin_monto"]:
        resp += f"  ({egr['pagos_fijos_sin_monto']} sin monto registrado, no incluidos - usa 'pago fijo X es QN' para fijarlo)\n"
    resp += f"- Cuentas por pagar que vencen el proximo mes: Q{egr['cuentas_por_pagar']:.2f} ({egr['cantidad_cuentas_pagar']} cuentas)\n"
    resp += f"- Compras estimadas por stock minimo: Q{egr['compras_estimadas']:.2f} ({egr['cantidad_productos_reabastecer']} productos)\n"
    resp += f"Total egresos: Q{egr['total']:.2f}\n\n"

    resp += "---\n"
    if f["positivo"]:
        resp += f"Flujo neto proyectado: +Q{f['flujo_neto']:.2f}\n"
        resp += f"Te sobraran aproximadamente Q{f['flujo_neto']:.2f} el proximo mes."
    else:
        resp += f"Flujo neto proyectado: -Q{abs(f['flujo_neto']):.2f}\n"
        resp += f"ALERTA: te faltaran aproximadamente Q{abs(f['flujo_neto']):.2f} para cubrir los egresos del proximo mes."
    return resp

register("contabilidad.flujo_proyectado_detalle", flujo_proyectado_detalle_tool)


def _resumen_iva(d):
    resp = f"Declaracion de IVA - {d['periodo']}\n\n"
    resp += f"Ventas totales: Q{d['total_ventas']:.2f} ({d['cantidad_ventas']} ventas)\n"
    resp += f"Debito fiscal (IVA cobrado, ventas x 12/112): Q{d['debito_fiscal']:.2f}\n"
    resp += f"Compras a proveedores registradas: Q{d['total_compras']:.2f} ({d['cantidad_ordenes_compra']} ordenes)\n"
    resp += f"Credito fiscal (IVA pagado, compras x 12%): Q{d['credito_fiscal']:.2f}\n"
    if d["cantidad_ordenes_compra"] == 0:
        resp += ("(Sin ordenes de compra con monto registrado este periodo - si tienes facturas "
                  "de proveedores fuera del sistema, sumalas manualmente al credito fiscal.)\n")
    resp += "---\n"
    if d["a_favor"]:
        resp += f"Tienes credito fiscal a favor de Q{abs(d['iva_a_pagar']):.2f} (no pagas IVA este mes, se traslada como remanente).\n"
    else:
        resp += f"IVA a pagar: Q{d['iva_a_pagar']:.2f}\n"
    resp += f"\n{d['plazo']}.\n"
    resp += "Este resumen es de referencia - verifica los montos contra tus facturas antes de declarar."
    return resp


def declaracion_iva_tool(data):
    from database.aronium import get_declaracion_iva
    from services.pdf_iva import generar_pdf_declaracion_iva
    d = get_declaracion_iva()
    filename = generar_pdf_declaracion_iva(d)
    resp = _resumen_iva(d)
    resp += f'\n\n<a href="/api/download/{filename}" download style="color:var(--accent)">⬇️ Descargar PDF de referencia</a>'
    return resp

register("contabilidad.declaracion_iva", declaracion_iva_tool)


def iva_rapido_tool(data):
    from database.aronium import get_declaracion_iva
    d = get_declaracion_iva()
    if d["a_favor"]:
        return (f"Para {d['periodo']}: tienes credito fiscal a favor de Q{abs(d['iva_a_pagar']):.2f} "
                f"(debito Q{d['debito_fiscal']:.2f} - credito Q{d['credito_fiscal']:.2f}), no pagas IVA este mes.")
    return (f"Para {d['periodo']}: debes Q{d['iva_a_pagar']:.2f} de IVA "
            f"(debito Q{d['debito_fiscal']:.2f} - credito Q{d['credito_fiscal']:.2f}). "
            "Escribe 'declaracion de IVA' para el resumen completo con PDF.")

register("contabilidad.iva_rapido", iva_rapido_tool)
