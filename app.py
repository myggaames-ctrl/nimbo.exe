"""
Nimbo AI Dashboard - Backend Flask
Conecta con OpenClaw y Ollama local
"""
import os
import subprocess
import json
import csv
import io
import re
import secrets
import unicodedata
import logging
from functools import wraps
from datetime import datetime
from dotenv import load_dotenv

# Se carga .env ANTES de importar cualquier modulo propio (openclaw_bridge,
# services/*, etc.) - varios de esos modulos leen os.getenv/os.environ.get a
# nivel de modulo (ej. openclaw_bridge.TOKEN), no dentro de una funcion, asi
# que si load_dotenv() corriera despues de esos imports (como pasaba antes)
# esas variables se resolverian vacias sin importar que .env las tuviera.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, '.env'))

from flask import Flask, render_template, request, jsonify, session, redirect, url_for, g
from werkzeug.exceptions import HTTPException
from flask_cors import CORS
from flask_login import (
    LoginManager, login_user, logout_user, login_required, current_user
)
from openclaw_bridge import send_to_openclaw, check_openclaw_status
from core.nimbo import send as nimbo_send
from services.inventory import InventoryService
from core.action_engine import ActionEngine
from core.tool_registry import list_tools
from services.skills_service import SkillsService
from services.mail_service import mail, enviar_reset_password
from services.ollama_service import llamar_ollama, analizar_ollama, ollama_disponible
from core.auth import (
    get_usuario_by_id, verify_login, existe_admin,
    listar_usuarios, crear_usuario, set_usuario_activo, cambiar_password_usuario,
    set_permisos_usuario, MODULOS_DISPONIBLES,
    get_admin_activo, crear_token_reset, validar_token_reset, usar_token_reset
)
from core.configuracion import get_config, set_config
from core.empresa import (
    init_master, get_empresa_by_slug, listar_empresas, get_db_path, crear_empresa
)
from database.nimbo_db import (
    get_diagnostico, diagnostico_completado, guardar_diagnostico, resetear_diagnostico,
    saltar_diagnostico, contar_productos
)
import tools

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("nimbo.setup")

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET_KEY", secrets.token_hex(32))
CORS(app)

# ── Correo (recuperación de contraseña) ────────────────────
app.config["MAIL_SERVER"] = os.environ.get("MAIL_SERVER", "smtp.gmail.com")
app.config["MAIL_PORT"] = int(os.environ.get("MAIL_PORT", 587))
app.config["MAIL_USE_TLS"] = True
app.config["MAIL_USERNAME"] = os.environ.get("MAIL_USERNAME", "")
app.config["MAIL_PASSWORD"] = os.environ.get("MAIL_PASSWORD", "")
app.config["MAIL_DEFAULT_SENDER"] = os.environ.get("MAIL_USERNAME", "")
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "")
mail.init_app(app)

# ── Multi-tenant: master DB (directorio de empresas) ───────
# Antes de esta migracion, aqui se llamaba a init_usuarios()/init_configuracion()
# y a los 22 init_* de database/nimbo_db.py, uno por uno, contra el unico
# nimbo.db global del proyecto. Ahora cada empresa tiene su propio nimbo.db
# (data/empresas/<slug>/nimbo.db) y esas tablas se crean recien cuando la
# empresa se registra (ver core/empresa.py crear_empresa() ->
# _inicializar_tablas_empresa(), que llama a esas mismas 22 funciones con el
# db_path de la empresa nueva) - al arrancar el servidor ya no hay un "nimbo.db"
# unico que inicializar de entrada, solo el maestro que sabe donde esta cada uno.
init_master()

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(user_id):
    return get_usuario_by_id(user_id)


def _es_ruta_api():
    """True si el request es a un endpoint /api/* (o /login o /select-empresa vistos
    desde fetch() - en la practica solo /api/* llama a esto). Estas rutas las
    consume JavaScript con fetch()+JSON.parse(), nunca un navegador navegando
    directo - por eso NUNCA deben responder con una redireccion (302) ni con una
    pagina HTML de error: JSON.parse() sobre HTML explota con
    "Unexpected token '<'" en vez de mostrar el error real (sesion vencida, sin
    empresa elegida, JSON invalido, etc.). Ver BUGS_Y_FIXES.md para el bug real
    que este chequeo corrige."""
    return request.path.startswith("/api/")


@login_manager.unauthorized_handler
def unauthorized():
    """Reemplaza el redirect-a-/login por defecto de Flask-Login para rutas
    /api/*. Sin esto, un fetch('/api/chat') con la sesion vencida (o con un
    cookie viejo de antes de la migracion multi-tenant, sin empresa_slug)
    terminaba siguiendo la cadena de redirects /login -> /select-empresa y
    devolviendo la pagina HTML de seleccion de empresa con status 200 - el
    frontend intentaba hacer JSON.parse() sobre ese HTML y tronaba."""
    if _es_ruta_api():
        return jsonify({"error": "Tu sesión expiró. Vuelve a iniciar sesión."}), 401
    return redirect(url_for("login", next=request.path))


# ── Onboarding (primera instalación) ───────────────────────
ONBOARDING_PASOS = (1, 2, 3, 4, 5, 6, 7)


def _onboarding_completado():
    return get_config("onboarding_completado") == "1"


def _onboarding_paso_actual():
    try:
        paso = int(get_config("onboarding_step", "1"))
    except (TypeError, ValueError):
        paso = 1
    return paso if paso in ONBOARDING_PASOS else 1


# ── Multi-tenant: contexto de empresa por request ──────────
# Se registra ANTES que require_setup (Flask corre los before_request en el
# orden en que se definen) para que _onboarding_completado()/get_config() de
# abajo ya lean de la empresa correcta (g.empresa_db) cuando require_setup los
# llama.
@app.before_request
def set_empresa_context():
    slug = session.get("empresa_slug")
    g.empresa_slug = None
    g.empresa_db = None
    if slug:
        db_path = get_db_path(slug)
        if db_path:
            g.empresa_slug = slug
            g.empresa_db = db_path
        else:
            # Slug presente en la sesion pero ya no existe en nimbo_master.db
            # (empresa borrada, o un slug de prueba viejo - bug real
            # encontrado 2026-07-19, ver BUGS_Y_FIXES.md): antes esto dejaba
            # g.empresa_slug con el valor viejo igual, asi que index()/login()
            # (que solo chequean "hay ALGUN slug en la sesion", no si ese slug
            # sigue existiendo) nunca redirigian a /select-empresa - el
            # usuario quedaba atascado en un /login para una empresa
            # fantasma, sin mas salida visible que el link "Ver demo del
            # onboarding", que lo mandaba a /setup/demo/1 en vez del flujo
            # multi-tenant normal. Se limpia la sesion aca mismo para que
            # todo lo que venga despues (index(), login(), require_setup())
            # trate esto exactamente igual que "nunca elegiste una empresa".
            session.pop("empresa_slug", None)
            session.pop("empresa_db", None)


ENDPOINTS_SIN_EMPRESA = (
    "select_empresa", "register_empresa", "api_empresa_crear", "static",
)


@app.before_request
def require_setup():
    if request.endpoint in ("setup", "setup_paso", "setup_demo", "setup_demo_paso"):
        return None
    if request.endpoint in ENDPOINTS_SIN_EMPRESA:
        return None
    if not g.empresa_slug:
        # Sin empresa elegida todavia - lo resuelve la ruta raiz (redirige a
        # /select-empresa) o el endpoint que corresponda, no este hook.
        return None
    # OJO: a diferencia de la version single-tenant, aqui ya NO se cachea el
    # resultado en una variable global (antes _setup_completado) - en
    # multi-tenant cada empresa tiene su propio onboarding_completado, asi que
    # el chequeo tiene que repetirse por request (una consulta SQL mas por
    # request, pero es la unica forma correcta de no confundir el "ya
    # completo" de una empresa con el de otra).
    if not _onboarding_completado():
        if _es_ruta_api():
            # Mismo motivo que unauthorized() arriba: una ruta /api/* nunca
            # debe devolver un redirect ni una pagina HTML - fetch()+JSON.parse()
            # tronaria igual que con la sesion vencida.
            return jsonify({"error": "Esta empresa todavía no completó la configuración inicial."}), 409
        return redirect(url_for("setup"))


@app.errorhandler(HTTPException)
def error_json_para_api(e):
    """Red de seguridad general para rutas /api/*: cualquier HTTPException que
    Werkzeug/Flask genere ANTES de que el codigo de la ruta corra (ej. 400 si
    request.json no logra parsear el body, 415 si el Content-Type no es
    application/json, 404, 405) se convierte a JSON en vez de la pagina HTML de
    error por defecto de Flask - mismo motivo que unauthorized_handler y el
    chequeo de require_setup de arriba: cualquier respuesta HTML para una ruta
    que el frontend consume con fetch()+JSON.parse() rompe con
    "Unexpected token '<'" en vez de mostrar el error real. Las rutas que no son
    /api/* (paginas normales del navegador) conservan las paginas de error HTML
    de siempre."""
    if _es_ruta_api():
        return jsonify({"error": e.description}), e.code
    return e


def admin_required(f):
    @wraps(f)
    @login_required
    def wrapper(*args, **kwargs):
        if not current_user.is_admin:
            return jsonify({"error": "Acceso restringido a administrador"}), 403
        return f(*args, **kwargs)
    return wrapper


def admin_o_tecnico_required(f):
    @wraps(f)
    @login_required
    def wrapper(*args, **kwargs):
        if not (current_user.is_admin or current_user.is_tecnico):
            return jsonify({"error": "Acceso restringido a administrador o tecnico"}), 403
        return f(*args, **kwargs)
    return wrapper

inventory = InventoryService()
engine = ActionEngine()
skills = SkillsService()


# ── Historial de chat en memoria ───────────────────────────
chat_history = []
nimbo_context = {}

# ══════════════════════════════════════════════════════════
# RUTAS PRINCIPALES
# ══════════════════════════════════════════════════════════

@app.route("/")
def index():
    # Se revisa la empresa ANTES que la sesion de usuario (a diferencia del
    # @login_required de antes de esta migracion): sin empresa elegida no
    # tiene sentido preguntar login, porque verify_login() no sabria contra
    # cual nimbo.db validar las credenciales.
    if not session.get("empresa_slug"):
        return redirect(url_for("select_empresa"))
    if not current_user.is_authenticated:
        return login_manager.unauthorized()
    # Diagnostico de negocio + tour: solo le importan al admin (las preguntas
    # son sobre el negocio y el dueño, y el tour resalta paneles admin-only -
    # ninguno de los dos tiene sentido para un empleado o tecnico). La
    # deteccion de "esta empresa ya tiene datos" se hace aca, al entrar al
    # dashboard (nunca en /diagnostico) - por diseño: si /diagnostico hiciera
    # su propio chequeo, una empresa con productos que llegara ahi por un
    # link viejo/directo alcanzaria a ver el formulario un instante antes de
    # que algo la sacara.
    tiene_productos = contar_productos() > 0
    if current_user.is_admin and not diagnostico_completado():
        if tiene_productos:
            # Empresa con datos existentes (ej. admin re-registrado, o el
            # diagnostico se reseteo manualmente desde Configuracion) - no
            # tiene sentido forzarla a contestar el cuestionario de negocio
            # nuevo. Se marca como saltada (mismo Completado=1 que
            # guardar_diagnostico(), sin inventar respuestas - ver
            # saltar_diagnostico()) y de paso se marca el tour como
            # completado, para que tampoco se le muestre: el tour solo tiene
            # sentido inmediatamente despues del diagnostico real de una
            # empresa nueva, no para una que ya esta operando.
            saltar_diagnostico()
            set_config("tour_completado", "1")
        else:
            return redirect(url_for("diagnostico"))
    return render_template(
        "index.html", rol=current_user.rol, nombre=current_user.nombre,
        negocio_nombre=get_config("negocio_nombre", "Mi Cocina"),
        tour_pendiente=current_user.is_admin and get_config("tour_completado") != "1",
        chat_tutorial_pendiente=get_config("chat_tutorial_mostrado") != "1",
        # Solo importa si el tutorial todavia esta pendiente (ver
        # mostrarTutorialChatSiCorresponde() en templates/index.html) - una
        # empresa sin ningun producto cargado ve la guia de "sube tu Excel o
        # agrega productos" en vez del tutorial normal de comandos.
        sin_inventario=not tiene_productos,
    )


@app.route("/setup")
def setup():
    if _onboarding_completado():
        return redirect(url_for("login"))
    return redirect(url_for("setup_paso", paso=_onboarding_paso_actual()))


def _avanzar_onboarding(siguiente):
    set_config("onboarding_step", str(siguiente))
    return redirect(url_for("setup_paso", paso=siguiente))


@app.route("/setup/<int:paso>", methods=["GET", "POST"])
def setup_paso(paso):
    global _setup_completado

    if _onboarding_completado():
        return redirect(url_for("login"))
    if paso not in ONBOARDING_PASOS:
        return redirect(url_for("setup"))

    paso_actual = _onboarding_paso_actual()
    if paso > paso_actual:
        return redirect(url_for("setup_paso", paso=paso_actual))

    error = None

    if request.method == "POST":
        saltar = request.form.get("accion") == "saltar"

        if paso == 1:
            negocio_nombre = request.form.get("negocio_nombre", "").strip()
            tipo_negocio = request.form.get("tipo_negocio", "").strip()
            pais = request.form.get("pais", "").strip()
            negocio_nit = request.form.get("negocio_nit", "").strip()
            negocio_telefono = request.form.get("negocio_telefono", "").strip()
            negocio_correo = request.form.get("negocio_correo", "").strip()
            negocio_direccion = request.form.get("negocio_direccion", "").strip()
            if not negocio_nombre or not tipo_negocio or not pais:
                error = "Nombre del negocio, tipo de negocio y pais son obligatorios."
            else:
                set_config("negocio_nombre", negocio_nombre)
                set_config("negocio_tipo", tipo_negocio)
                set_config("negocio_pais", pais)
                set_config("negocio_nit", negocio_nit)
                set_config("negocio_telefono", negocio_telefono)
                set_config("negocio_correo", negocio_correo)
                set_config("negocio_direccion", negocio_direccion)
                return _avanzar_onboarding(2)

        elif paso == 2:
            if existe_admin():
                # La cuenta admin ya se creo en un intento anterior de este paso.
                # OJO: antes esto revisaba existe_algun_usuario() (cualquier fila
                # en Usuario), lo que hacia que este paso se saltara crear_usuario()
                # en silencio si ya habia un empleado/tecnico de un intento de
                # onboarding anterior, aunque nunca se hubiera creado un admin -
                # ver existe_admin() en core/auth.py.
                logger.warning(
                    "setup paso 2: ya existe un admin, se omite crear_usuario() "
                    "(dueno_nombre=%r, username=%r)",
                    request.form.get("dueno_nombre", ""), request.form.get("username", "")
                )
                return _avanzar_onboarding(3)
            dueno_nombre = request.form.get("dueno_nombre", "").strip()
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            email = request.form.get("email", "").strip()
            if not dueno_nombre or not username or not password:
                error = "Nombre, usuario y contraseña son obligatorios."
            elif len(password) < 6:
                error = "La contraseña debe tener al menos 6 caracteres."
            else:
                try:
                    crear_usuario(username, password, dueno_nombre, "admin")
                    set_config("dueno_nombre", dueno_nombre)
                    set_config("admin_email", email)
                    logger.info(
                        "setup paso 2: admin creado OK (username=%r, dueno_nombre=%r)",
                        username, dueno_nombre
                    )
                    return _avanzar_onboarding(3)
                except Exception as e:
                    logger.exception(
                        "setup paso 2: fallo crear_usuario() (username=%r, dueno_nombre=%r)",
                        username, dueno_nombre
                    )
                    error = str(e)

        elif paso == 3:
            moneda = request.form.get("moneda", "").strip() or "Q"
            cobra_iva = "si" if request.form.get("cobra_iva") == "si" else "no"
            usa_pos = "si" if request.form.get("usa_pos") == "si" else "no"
            isr_trimestral = "si" if request.form.get("isr_trimestral") == "si" else "no"
            comision_pos = request.form.get("comision_pos", "").strip() or "6"
            if usa_pos == "si":
                try:
                    comision_pos = str(round(float(comision_pos), 2))
                except ValueError:
                    comision_pos = "6"
            set_config("moneda", moneda)
            set_config("cobra_iva", cobra_iva)
            set_config("iva_porcentaje", "12" if cobra_iva == "si" else "0")
            set_config("usa_pos", usa_pos)
            set_config("comision_pos", comision_pos if usa_pos == "si" else "0")
            set_config("isr_trimestral", isr_trimestral)
            set_config("isr_porcentaje", "5" if isr_trimestral == "si" else "0")
            return _avanzar_onboarding(4)

        elif paso == 4:
            apertura = request.form.get("horario_apertura", "").strip() or "08:00"
            cierre = request.form.get("horario_cierre", "").strip() or "18:00"
            dias = request.form.getlist("dias_laborales")
            set_config("horario_apertura", apertura)
            set_config("horario_cierre", cierre)
            set_config("dias_laborales", ",".join(dias))
            return _avanzar_onboarding(5)

        elif paso == 5:
            if saltar:
                return _avanzar_onboarding(6)
            fondo = request.form.get("caja_fondo_inicial", "").strip()
            if not fondo:
                return _avanzar_onboarding(6)
            try:
                monto = round(float(fondo), 2)
                if monto < 0:
                    raise ValueError()
                set_config("caja_fondo_inicial", str(monto))
                # Reutiliza el mismo mecanismo que el cierre de turno usa para
                # sugerir el fondo del dia siguiente (ver database.nimbo_db.
                # get_fondo_sugerido) - asi el primer "abrir caja" del negocio
                # ya llega prellenado con este monto, en vez de pedirlo nuevo.
                set_config("caja_fondo_sugerido", str(monto))
                return _avanzar_onboarding(6)
            except ValueError:
                error = "El fondo de caja debe ser un numero valido, o salta este paso."

        elif paso == 6:
            if saltar:
                return _avanzar_onboarding(7)
            emp_nombre = request.form.get("empleado_nombre", "").strip()
            emp_username = request.form.get("empleado_username", "").strip()
            emp_password = request.form.get("empleado_password", "")
            emp_rol = request.form.get("empleado_rol", "empleado").strip()
            emp_telefono = request.form.get("empleado_telefono", "").strip()
            if not emp_nombre or not emp_username or not emp_password:
                error = "Escribe nombre, usuario y contraseña, o salta este paso."
            elif len(emp_password) < 6:
                error = "La contraseña debe tener al menos 6 caracteres."
            elif emp_rol not in ("empleado", "tecnico"):
                error = "Rol invalido."
            else:
                try:
                    crear_usuario(emp_username, emp_password, emp_nombre, emp_rol, emp_telefono)
                    set_config("onboarding_empleado_nombre", emp_nombre)
                    return _avanzar_onboarding(7)
                except Exception as e:
                    error = str(e)

        elif paso == 7:
            set_config("onboarding_completado", "1")
            _setup_completado = True
            return redirect(url_for("login"))

    resumen = None
    if paso == 7:
        resumen = {
            "negocio_nombre": get_config("negocio_nombre", ""),
            "negocio_tipo": get_config("negocio_tipo", ""),
            "negocio_pais": get_config("negocio_pais", ""),
            "negocio_nit": get_config("negocio_nit", ""),
            "negocio_telefono": get_config("negocio_telefono", ""),
            "negocio_correo": get_config("negocio_correo", ""),
            "negocio_direccion": get_config("negocio_direccion", ""),
            "dueno_nombre": get_config("dueno_nombre", ""),
            "moneda": get_config("moneda", "Q"),
            "cobra_iva": get_config("cobra_iva", "no"),
            "usa_pos": get_config("usa_pos", "no"),
            "comision_pos": get_config("comision_pos", "0"),
            "isr_trimestral": get_config("isr_trimestral", "no"),
            "horario_apertura": get_config("horario_apertura", ""),
            "horario_cierre": get_config("horario_cierre", ""),
            "dias_laborales": get_config("dias_laborales", ""),
            "caja_fondo_inicial": get_config("caja_fondo_inicial", ""),
            "empleado_nombre": get_config("onboarding_empleado_nombre", ""),
        }

    return render_template("setup.html", paso=paso, total_pasos=len(ONBOARDING_PASOS),
                            error=error, resumen=resumen, demo=False)


# ── Demo del onboarding (sin login, no persiste nada) ──────
# Accesible desde el boton "Ver demo del onboarding" en /login. Reutiliza el
# mismo template que /setup/<paso>, pero con demo=True: los formularios pasan
# a GET (no hay ruta de procesamiento que valide ni llame set_config/
# crear_usuario) y el resumen del paso 7 usa datos de ejemplo fijos en vez de
# leer Configuracion - esta ruta no requiere sesion, asi que jamas debe
# mostrar los datos reales del negocio.
DEMO_RESUMEN_EJEMPLO = {
    "negocio_nombre": "Mi Cocina Electrodomésticos",
    "negocio_tipo": "Electrodomésticos y repuestos",
    "negocio_pais": "Guatemala",
    "negocio_nit": "1234567-8",
    "negocio_telefono": "5555-5555",
    "negocio_correo": "contacto@ejemplo.com",
    "negocio_direccion": "5a avenida 10-20 zona 1, Guatemala",
    "dueno_nombre": "Nombre del administrador",
    "moneda": "Q",
    "cobra_iva": "si",
    "usa_pos": "si",
    "comision_pos": "6",
    "isr_trimestral": "si",
    "horario_apertura": "08:00",
    "horario_cierre": "18:00",
    "dias_laborales": "Lunes,Martes,Miércoles,Jueves,Viernes,Sábado",
    "caja_fondo_inicial": "400",
    "empleado_nombre": "Nombre del empleado",
}


@app.route("/setup/demo")
def setup_demo():
    return redirect(url_for("setup_demo_paso", paso=1))


@app.route("/setup/demo/<int:paso>")
def setup_demo_paso(paso):
    if paso not in ONBOARDING_PASOS:
        return redirect(url_for("setup_demo"))
    resumen = DEMO_RESUMEN_EJEMPLO if paso == 7 else None
    return render_template("setup.html", paso=paso, total_pasos=len(ONBOARDING_PASOS),
                            error=None, resumen=resumen, demo=True)


@app.route("/login", methods=["GET", "POST"])
def login():
    if not session.get("empresa_slug"):
        return redirect(url_for("select_empresa"))

    if current_user.is_authenticated:
        return redirect(url_for("index"))

    error = None
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = verify_login(username, password)
        if user:
            login_user(user)
            # Se re-guarda la empresa en session al loguear (ademas de ya
            # estar ahi desde /select-empresa) para dejar tambien empresa_db
            # cacheado en session, tal como pide el plan de migracion.
            empresa = get_empresa_by_slug(session["empresa_slug"])
            if empresa:
                session["empresa_slug"] = empresa["Slug"]
                session["empresa_db"] = empresa["DbPath"]
            next_page = request.args.get("next")
            return redirect(next_page or url_for("index"))
        error = "Usuario o contraseña incorrectos"

    return render_template("login.html", error=error)


@app.route("/select-empresa", methods=["GET", "POST"])
def select_empresa():
    if request.method == "POST":
        slug = request.form.get("slug", "").strip()
        empresa = get_empresa_by_slug(slug)
        if not empresa:
            return render_template("select_empresa.html", empresas=listar_empresas(),
                                    error="Empresa no encontrada.")
        session["empresa_slug"] = empresa["Slug"]
        session["empresa_db"] = empresa["DbPath"]
        return redirect(url_for("login"))
    return render_template("select_empresa.html", empresas=listar_empresas(), error=None)


@app.route("/register-empresa")
def register_empresa():
    return render_template("register_empresa.html")


@app.route("/api/empresa/crear", methods=["POST"])
def api_empresa_crear():
    data = request.json or {}
    nombre = (data.get("nombre") or "").strip()
    admin_username = (data.get("admin_username") or "").strip()
    admin_password = data.get("admin_password") or ""
    admin_nombre = (data.get("admin_nombre") or "").strip()

    if not nombre:
        return jsonify({"error": "El nombre de la empresa es obligatorio"}), 400
    if not admin_username or not admin_password or not admin_nombre:
        return jsonify({"error": "Faltan datos del administrador"}), 400
    if len(admin_password) < 6:
        return jsonify({"error": "La contraseña debe tener al menos 6 caracteres"}), 400

    try:
        empresa = crear_empresa(nombre)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    db_path_absoluto = os.path.join(BASE_DIR, empresa["db_path"])
    try:
        crear_usuario(admin_username, admin_password, admin_nombre, "admin",
                       db_path=db_path_absoluto)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400

    # A diferencia de Mi Cocina (que paso por los 7 pasos del wizard viejo,
    # incluyendo datos financieros/horario/fondo de caja), una empresa creada
    # por este registro rapido de 2 pasos no pasa por esos pasos - se marca
    # onboarding_completado directo para que no quede atrapada pidiendo el
    # wizard completo la primera vez que alguien intente loguearse. El nombre
    # del negocio si se guarda, para que el sidebar del dashboard lo muestre.
    set_config("negocio_nombre", nombre, db_path=db_path_absoluto)
    set_config("onboarding_completado", "1", db_path=db_path_absoluto)

    session["empresa_slug"] = empresa["slug"]
    session["empresa_db"] = empresa["db_path"]

    return jsonify({"ok": True, "slug": empresa["slug"]})


@app.route("/diagnostico", methods=["GET"])
@login_required
def diagnostico():
    # Solo el admin lo completa (las preguntas son sobre el negocio y el
    # dueño) y solo si todavia no lo completo - si un empleado llega aca por
    # URL directa, o un admin que ya lo completo, se lo manda de vuelta al
    # dashboard en vez de un 403 (esta es una pagina normal del navegador,
    # no una ruta /api/*, asi que un redirect es lo esperado aca). Esta ruta
    # NO revisa contar_productos() - esa deteccion (saltar automaticamente el
    # diagnostico de una empresa con datos existentes) se hace en index() al
    # entrar al dashboard, no aca. Si esta funcion se alcanza de verdad es
    # porque index() ya determino que la empresa no tiene productos todavia,
    # asi que el formulario siempre se muestra completo, sin boton de saltar.
    if not current_user.is_admin or diagnostico_completado():
        return redirect(url_for("index"))
    return render_template("diagnostico.html", nombre=current_user.nombre)


@app.route("/api/diagnostico/guardar", methods=["POST"])
@admin_required
def api_diagnostico_guardar():
    data = request.json or {}
    resultado = guardar_diagnostico(data)
    return jsonify({"ok": True, **resultado})


@app.route("/api/diagnostico/resumen", methods=["GET"])
@admin_required
def api_diagnostico_resumen():
    """Para el panel Configuracion: fecha + areas de oportunidad del ultimo
    diagnostico completado, o completado=False si nunca se hizo uno."""
    d = get_diagnostico()
    if not d or not d.get("Completado"):
        return jsonify({"completado": False})
    return jsonify({
        "completado": True,
        "fecha_completado": d.get("FechaCompletado"),
        "resumen": d.get("ResumenDiagnostico"),
        "areas_oportunidad": d.get("AreasOportunidad"),
        "primeros_tres_pasos": d.get("PrimerosTresPasos"),
    })


@app.route("/api/diagnostico/repetir", methods=["POST"])
@admin_required
def api_diagnostico_repetir():
    """Boton 'Repetir diagnostico de negocio' de Configuracion: solo marca
    Completado=0 (no borra las respuestas anteriores, ver
    resetear_diagnostico) - el frontend redirige a /diagnostico despues de
    esto, que vuelve a mostrar el formulario vacio."""
    resetear_diagnostico()
    return jsonify({"ok": True})


@app.route("/api/tour/completar", methods=["POST"])
@admin_required
def api_tour_completar():
    set_config("tour_completado", "1")
    return jsonify({"ok": True})


@app.route("/api/tour/resetear", methods=["POST"])
@admin_required
def api_tour_resetear():
    """Boton 'Ver tour del dashboard' de Configuracion - el frontend recarga
    la pagina despues de esto, y el tour arranca de nuevo porque
    tour_pendiente vuelve a ser True en index()."""
    set_config("tour_completado", "0")
    return jsonify({"ok": True})


@app.route("/api/chat/tutorial-mostrado", methods=["POST"])
@login_required
def api_chat_tutorial_mostrado():
    """Cualquier rol puede disparar esto (no solo admin) - el mensaje de
    tutorial del chat lo ve quien sea que abra el chat primero, no solo el
    admin que hizo el diagnostico."""
    set_config("chat_tutorial_mostrado", "1")
    return jsonify({"ok": True})


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    mensaje = None
    if request.method == "POST":
        email = request.form.get("email", "").strip()
        mensaje = "Si el correo es válido, te enviamos un enlace de recuperación."

        if ADMIN_EMAIL and email.lower() == ADMIN_EMAIL.lower():
            admin_user = get_admin_activo()
            if admin_user:
                token = crear_token_reset(admin_user.id)
                link = url_for("reset_password", token=token, _external=True)
                try:
                    enviar_reset_password(ADMIN_EMAIL, link)
                except Exception as e:
                    print("Error enviando correo de recuperación:", e)

        return render_template("forgot_password.html", mensaje=mensaje)

    return render_template("forgot_password.html", mensaje=mensaje)


@app.route("/reset-password", methods=["GET", "POST"])
def reset_password():
    token = request.args.get("token", "") or request.form.get("token", "")
    user = validar_token_reset(token) if token else None

    if not token or not user:
        return render_template("reset_password.html", token_valido=False)

    if request.method == "POST":
        password = request.form.get("password", "")
        password2 = request.form.get("password2", "")
        if not password or len(password) < 6:
            return render_template("reset_password.html", token_valido=True, token=token,
                                    error="La contraseña debe tener al menos 6 caracteres.")
        if password != password2:
            return render_template("reset_password.html", token_valido=True, token=token,
                                    error="Las contraseñas no coinciden.")
        if usar_token_reset(token, password):
            return redirect(url_for("login", reset="ok"))
        return render_template("reset_password.html", token_valido=False)

    return render_template("reset_password.html", token_valido=True, token=token)


# ══════════════════════════════════════════════════════════
# USUARIOS (solo admin)
# ══════════════════════════════════════════════════════════

@app.route("/api/usuarios", methods=["GET"])
@admin_required
def api_usuarios_list():
    return jsonify(listar_usuarios())


@app.route("/api/usuarios/modulos", methods=["GET"])
@admin_required
def api_usuarios_modulos():
    return jsonify(list(MODULOS_DISPONIBLES))


@app.route("/api/usuarios", methods=["POST"])
@admin_required
def api_usuarios_create():
    data = request.json
    try:
        crear_usuario(
            data.get("username", ""),
            data.get("password", ""),
            data.get("nombre", ""),
            data.get("rol", "empleado"),
            data.get("telefono", ""),
            permisos=data.get("permisos")
        )
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@app.route("/api/usuarios/<int:uid>/toggle", methods=["POST"])
@admin_required
def api_usuarios_toggle(uid):
    data = request.json or {}
    set_usuario_activo(uid, data.get("activo", True))
    return jsonify({"ok": True})


@app.route("/api/usuarios/<int:uid>/password", methods=["PUT"])
@admin_required
def api_usuarios_password(uid):
    data = request.json or {}
    try:
        cambiar_password_usuario(uid, data.get("password", ""))
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 400


@app.route("/api/usuarios/<int:uid>/permisos", methods=["PUT"])
@admin_required
def api_usuarios_permisos(uid):
    data = request.json or {}
    try:
        permisos = set_permisos_usuario(uid, data.get("permisos", []))
        return jsonify({"ok": True, "permisos": permisos})
    except Exception as e:
        return jsonify({"error": str(e)}), 400


# ── Configuracion del negocio (panel "Configuracion" del dashboard) ────────
# Mismos datos que se recopilan en /setup (pasos 1, 3, 4 y 5), pero editables
# despues sin tener que resetear onboarding_completado. No incluye la cuenta
# admin (paso 2) ni el primer empleado (paso 6) - esos ya se gestionan desde
# el panel "Usuarios".
CONFIG_NEGOCIO_KEYS = (
    "negocio_nombre", "negocio_tipo", "negocio_pais", "negocio_nit",
    "negocio_telefono", "negocio_correo", "negocio_direccion",
    "moneda", "cobra_iva", "usa_pos", "comision_pos", "isr_trimestral",
    "horario_apertura", "horario_cierre", "dias_laborales",
    "caja_fondo_inicial",
)


@app.route("/api/configuracion/negocio", methods=["GET"])
@admin_required
def api_configuracion_negocio_get():
    return jsonify({k: get_config(k, "") for k in CONFIG_NEGOCIO_KEYS})


@app.route("/api/configuracion/negocio", methods=["PUT"])
@admin_required
def api_configuracion_negocio_update():
    data = request.json or {}

    negocio_nombre = (data.get("negocio_nombre") or "").strip()
    tipo_negocio = (data.get("negocio_tipo") or "").strip()
    pais = (data.get("negocio_pais") or "").strip()
    if not negocio_nombre or not tipo_negocio or not pais:
        return jsonify({"error": "Nombre del negocio, tipo de negocio y pais son obligatorios."}), 400

    cobra_iva = "si" if data.get("cobra_iva") == "si" else "no"
    usa_pos = "si" if data.get("usa_pos") == "si" else "no"
    isr_trimestral = "si" if data.get("isr_trimestral") == "si" else "no"
    comision_pos = str(data.get("comision_pos") or "6").strip()
    if usa_pos == "si":
        try:
            comision_pos = str(round(float(comision_pos), 2))
        except (TypeError, ValueError):
            comision_pos = "6"

    dias = data.get("dias_laborales")
    if isinstance(dias, list):
        dias_laborales = ",".join(dias)
    else:
        dias_laborales = (dias or "").strip()

    set_config("negocio_nombre", negocio_nombre)
    set_config("negocio_tipo", tipo_negocio)
    set_config("negocio_pais", pais)
    set_config("negocio_nit", (data.get("negocio_nit") or "").strip())
    set_config("negocio_telefono", (data.get("negocio_telefono") or "").strip())
    set_config("negocio_correo", (data.get("negocio_correo") or "").strip())
    set_config("negocio_direccion", (data.get("negocio_direccion") or "").strip())
    set_config("moneda", (data.get("moneda") or "Q").strip() or "Q")
    set_config("cobra_iva", cobra_iva)
    set_config("iva_porcentaje", "12" if cobra_iva == "si" else "0")
    set_config("usa_pos", usa_pos)
    set_config("comision_pos", comision_pos if usa_pos == "si" else "0")
    set_config("isr_trimestral", isr_trimestral)
    set_config("isr_porcentaje", "5" if isr_trimestral == "si" else "0")
    set_config("horario_apertura", (data.get("horario_apertura") or "").strip())
    set_config("horario_cierre", (data.get("horario_cierre") or "").strip())
    set_config("dias_laborales", dias_laborales)

    fondo = (data.get("caja_fondo_inicial") or "").strip()
    if fondo:
        try:
            monto = round(float(fondo), 2)
            if monto < 0:
                raise ValueError()
        except ValueError:
            return jsonify({"error": "El fondo de caja debe ser un numero valido."}), 400
        set_config("caja_fondo_inicial", str(monto))
        # Mismo mecanismo que usa el cierre de turno para sugerir el fondo del
        # siguiente turno (ver database.nimbo_db.cerrar_turno/get_fondo_sugerido).
        set_config("caja_fondo_sugerido", str(monto))
    else:
        set_config("caja_fondo_inicial", "")

    return jsonify({"ok": True})


# ══════════════════════════════════════════════════════════
# CHAT CON NIMBO (OLLAMA) / OPENCLAW
# ══════════════════════════════════════════════════════════

@app.route("/api/chat", methods=["POST"])
@login_required
def chat():

    data = request.json
    mensaje = data.get("message", "").strip()

    if not mensaje:
        return jsonify({"error": "Mensaje vacío"}), 400

    # ==================================================
    # Resolver usando NimboCore (herramientas + fallback final a Ollama,
    # ver core/nimbo.py send() - siempre devuelve una respuesta no vacia,
    # asi que no hace falta un segundo intento con otro modelo aqui).
    # ==================================================
    from core.nimbo import send as nimbo_send
    ctx = session.get("nimbo_ctx", {})
    try:
        respuesta_tool, ctx = nimbo_send(mensaje, context=ctx, role=current_user.rol)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    session["nimbo_ctx"] = ctx

    is_html = "<a " in str(respuesta_tool)
    chat_history.append({"role": "user", "content": mensaje, "time": datetime.now().strftime("%H:%M")})
    chat_history.append({"role": "assistant", "content": respuesta_tool, "time": datetime.now().strftime("%H:%M")})
    from database.nimbo_db import guardar_mensaje_chat
    guardar_mensaje_chat("user", mensaje)
    guardar_mensaje_chat("assistant", str(respuesta_tool))
    return jsonify({"response": respuesta_tool, "source": "tool", "html": is_html})


@app.route("/api/chat/history", methods=["GET"])
@login_required
def get_history():
    return jsonify(chat_history)


@app.route("/api/chat/clear", methods=["POST"])
@login_required
def clear_history():
    chat_history.clear()
    return jsonify({"ok": True})


# ══════════════════════════════════════════════════════════
# INVENTARIO (catalogo propio de Nimbo)
# ══════════════════════════════════════════════════════════
# Hasta la remocion completa de la integracion con Aronium (ver BUGS_Y_FIXES.md)
# estas rutas leian pos.db (Aronium) directo, adivinando nombres de tabla
# genericos ("items"/"products"/etc). Ahora usan las funciones propias de
# database/nimbo_db.py, igual que el resto del dashboard. Existia ademas una
# ruta duplicada /api/inventory/products (inventory_products(), mas abajo en
# el archivo) que llamaba a un metodo inexistente (inventory.products()) y
# por eso nunca se ejecutaba en la practica - Werkzeug siempre despachaba a
# la version de aca (la primera registrada); se elimino esa duplicada al
# hacer esta limpieza.

@app.route("/api/inventory/products", methods=["GET"])
@login_required
def get_products():
    try:
        from database.nimbo_db import get_products as db_get_products
        search = request.args.get("search", "")
        productos = db_get_products(search=search)
        return jsonify({
            "products": productos,
            "columns": ["Id", "Name", "Code", "Price", "Cost", "Stock"],
            "table": "Producto",
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/inventory/stats", methods=["GET"])
@login_required
def get_stats():
    try:
        from database.nimbo_db import contar_productos, get_sales_month
        return jsonify({
            "total_products": contar_productos(),
            "total_sales": get_sales_month()["Cantidad"],
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/inventory/query", methods=["POST"])
@admin_required
def custom_query():
    """Consulta SQL personalizada (solo SELECT) contra nimbo.db"""
    data = request.json
    sql = data.get("sql", "").strip()
    if not sql.lower().startswith("select"):
        return jsonify({"error": "Solo se permiten consultas SELECT"}), 400
    try:
        from database.nimbo_db import nimbo_connect
        conn = nimbo_connect()
        cur = conn.cursor()
        cur.execute(sql)
        rows = [dict(r) for r in cur.fetchall()]
        cols = [d[0] for d in cur.description]
        conn.close()
        return jsonify({"rows": rows, "columns": cols})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/inventory/analyze", methods=["POST"])
@login_required
def analyze_inventory():
    """Análisis con Ollama local sobre el inventario"""
    data = request.json
    context = data.get("context", "")
    question = data.get("question", "Dame un análisis del inventario")
    try:
        prompt = f"""Eres un analista de negocios para Mi Cocina Electrodomésticos en Guatemala.
Datos del inventario:
{context}

Pregunta: {question}
Responde en español con análisis claro y recomendaciones prácticas."""
        texto = analizar_ollama(prompt)
        if not texto:
            return jsonify({"error": "Ollama no esta disponible. Verifica que este corriendo."}), 500
        return jsonify({"analysis": texto})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ══════════════════════════════════════════════════════════
# GENERADOR DE CONTENIDO
# ══════════════════════════════════════════════════════════

@app.route("/api/content/generate", methods=["POST"])
@admin_required
def generate_content():
    data = request.json
    tipo = data.get("type", "post")
    producto = data.get("product", "")
    tono = data.get("tone", "profesional")

    prompts = {
        "post": f"""Crea un post para redes sociales para Mi Cocina Electrodomésticos Guatemala.
Producto: {producto}
Tono: {tono}
Incluye: texto principal, emojis, call to action y hashtags.
Formato: breve, para Instagram/Facebook.""",

        "script": f"""Crea un script para video/reel de 30-45 segundos para Mi Cocina Electrodomésticos Guatemala.
Producto: {producto}
Tono: {tono}
Incluye: gancho inicial, beneficios clave, llamada a la acción.
Formato: guión con indicaciones de escena.""",

        "description": f"""Crea una descripción de producto para Mi Cocina Electrodomésticos Guatemala.
Producto: {producto}
Tono: {tono}
Incluye: características principales, beneficios, precio (si aplica: "consultar precio").
Formato: descripción atractiva para catálogo/tienda online.""",

        "hashtags": f"""Genera 20 hashtags relevantes para Mi Cocina Electrodomésticos Guatemala.
Producto/tema: {producto}
Mix de: hashtags populares, nicho, y locales de Guatemala."""
    }

    prompt = prompts.get(tipo, prompts["post"])

    try:
        texto = llamar_ollama(prompt)
        if not texto:
            return jsonify({"error": "Ollama no esta disponible. Verifica que este corriendo."}), 500
        return jsonify({"content": texto, "type": tipo})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ══════════════════════════════════════════════════════════
# TAREAS / NOTAS
# ══════════════════════════════════════════════════════════

tasks = []

@app.route("/api/tasks", methods=["GET"])
@admin_required
def get_tasks():
    return jsonify(tasks)


@app.route("/api/tasks", methods=["POST"])
@admin_required
def add_task():
    data = request.json
    task = {
        "id": len(tasks) + 1,
        "text": data.get("text", ""),
        "done": False,
        "created": datetime.now().strftime("%d/%m %H:%M")
    }
    tasks.append(task)
    return jsonify(task)


@app.route("/api/tasks/<int:task_id>", methods=["PATCH"])
@admin_required
def update_task(task_id):
    for t in tasks:
        if t["id"] == task_id:
            t["done"] = not t["done"]
            return jsonify(t)
    return jsonify({"error": "No encontrado"}), 404


@app.route("/api/tasks/<int:task_id>", methods=["DELETE"])
@admin_required
def delete_task(task_id):
    global tasks
    tasks = [t for t in tasks if t["id"] != task_id]
    return jsonify({"ok": True})


# ══════════════════════════════════════════════════════════
# SISTEMA
# ══════════════════════════════════════════════════════════

@app.route("/api/status", methods=["GET"])
@login_required
def status():
    return jsonify({
        "ollama": ollama_disponible(),
        "time": datetime.now().strftime("%d/%m/%Y %H:%M")
    })
@app.route("/api/openclaw/chat", methods=["POST"])
@login_required
def openclaw_chat():

    data = request.json
    message = data.get("message", "").strip()

    if not message:
        return jsonify({
            "error": "Mensaje vacío"
        }), 400

    # ==================================================
    # 1. Intentar resolver usando herramientas
    # ==================================================

    from core.nimbo import send as nimbo_send_ctx
    ctx = session.get("nimbo_ctx", {})
    response, ctx = nimbo_send_ctx(message, context=ctx, role=current_user.rol)
    session["nimbo_ctx"] = ctx
    from database.nimbo_db import guardar_mensaje_chat
    guardar_mensaje_chat("user", message)
    guardar_mensaje_chat("assistant", str(response))

    chat_history.append({
        "role": "user",
        "content": message,
        "time": datetime.now().strftime("%H:%M")
    })

    chat_history.append({
        "role": "assistant",
        "content": response,
        "time": datetime.now().strftime("%H:%M")
    })

    if not response:
        return jsonify({
        "error": "OpenClaw no devolvió ninguna respuesta."
        }), 500

    is_html = "<a " in str(response) or "<button" in str(response)
    return jsonify({
        "response": response,
        "source": "openclaw",
        "html": is_html
    })

@app.route("/api/inventory/product-columns", methods=["GET"])
@login_required
def product_columns():

    try:
        return jsonify(inventory.product_columns())
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/inventory/stock-columns", methods=["GET"])
@login_required
def stock_columns():

    try:
        return jsonify(inventory.stock_columns())
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    
@app.route("/api/skills/check", methods=["GET"])
@admin_required
def skills_check():

    try:

        return jsonify(

            skills.status()

        )

    except Exception as e:

        return jsonify({

            "error": str(e)

        }), 500
      
@app.route("/api/inventory/search")
@login_required
def inventory_search():

    barcode = request.args.get("barcode", "").strip()

    product = inventory.search_barcode(barcode)

    if product:
        return jsonify(product)

    return jsonify({
        "error": "Producto no encontrado"
    }), 404

@app.route("/api/tool/inventory/search", methods=["GET"])
@login_required
def tool_inventory_search():

    text = request.args.get("q", "")

    return jsonify({
        "response": search_product(text)
    })

@app.route("/pos")
@login_required
def pos():

    return render_template("pos.html")

@app.route("/api/tools")
@admin_required
def api_tools():
    return jsonify({
        "tools": list_tools()
    })


@app.route("/api/db/products")
@login_required
def db_products():
    try:
        from database.nimbo_db import get_products
        search = request.args.get("search", "")
        return jsonify(get_products(search=search))
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/db/stats")
@login_required
def db_stats():
    try:
        from database.nimbo_db import get_sales_today, get_sales_month, get_low_stock, get_top_products
        return jsonify({
            "sales_today": get_sales_today(),
            "sales_month": get_sales_month(),
            "low_stock": get_low_stock(),
            "top_products": get_top_products()
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/db/customers")
@login_required
def db_customers():
    try:
        from database.nimbo_db import get_customers
        search = request.args.get("search", "")
        return jsonify(get_customers(search=search))
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/db/products/add", methods=["POST"])
@login_required
def db_add_product():
    from database.nimbo_db import nimbo_connect
    data = request.json
    try:
        conn = nimbo_connect()
        cur = conn.cursor()
        cur.execute("INSERT INTO Producto (Nombre, Codigo, Precio, Costo, Activo) VALUES (?, ?, ?, ?, 1)",
            (data["Name"], data["Code"], data["Price"], data["Cost"]))
        conn.commit()
        conn.close()
        return jsonify({"ok": True, "id": cur.lastrowid})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/db/products/<int:pid>", methods=["PUT"])
@login_required
def db_edit_product(pid):
    from database.nimbo_db import nimbo_connect
    data = request.json
    try:
        conn = nimbo_connect()
        cur = conn.cursor()
        cur.execute("UPDATE Producto SET Nombre=?, Codigo=?, Precio=?, Costo=? WHERE Id=?",
            (data["Name"], data["Code"], data["Price"], data["Cost"], pid))
        conn.commit()
        conn.close()
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/db/products/<int:pid>", methods=["DELETE"])
@login_required
def db_delete_product(pid):
    from database.nimbo_db import nimbo_connect
    try:
        conn = nimbo_connect()
        cur = conn.cursor()
        cur.execute("UPDATE Producto SET Activo=0 WHERE Id=?", (pid,))
        conn.commit()
        conn.close()
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── Importacion masiva de productos (Excel/CSV) ─────────────────────────
# Columnas reconocidas (con o sin acentos/mayusculas/espacios, ver
# _normalizar_header): Nombre y Precio son obligatorias, Costo/Codigo/Stock/
# StockMinimo/Categoria/Descripcion/Unidad son opcionales. Ver CLAUDE.md
# seccion "Importacion masiva de productos" para el flujo completo
# (deteccion de duplicados, generacion automatica de codigo de barras,
# preview + mapeo manual de columnas).
COLUMNAS_IMPORTACION_PRODUCTOS = {
    "nombre": "nombre", "precio": "precio", "costo": "costo",
    "codigo": "codigo", "stock": "stock", "stockminimo": "stock_minimo",
    "categoria": "categoria", "descripcion": "descripcion", "unidad": "unidad",
}

# Etiquetas en espanol para el selector de mapeo manual del modal de
# importacion (templates/index.html) - mismo orden en que conviene
# mostrarlas, campo interno -> texto visible.
CAMPOS_IMPORTACION_LABELS = {
    "nombre": "Nombre", "precio": "Precio", "costo": "Costo",
    "codigo": "Código", "stock": "Stock", "stock_minimo": "Stock mínimo",
    "categoria": "Categoría", "descripcion": "Descripción", "unidad": "Unidad",
}


def _normalizar_header(h):
    h = unicodedata.normalize("NFD", str(h or "")).encode("ascii", "ignore").decode("utf-8")
    return re.sub(r"[^a-z0-9]", "", h.lower())


def _parsear_numero_importado(valor):
    if valor is None:
        return None
    texto = str(valor).strip().replace("Q", "").replace("q", "").replace(",", "")
    if texto == "":
        return None
    return float(texto)


def _leer_filas_csv_productos(archivo):
    texto = archivo.stream.read().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(texto)))


def _leer_filas_excel_productos(archivo):
    import openpyxl
    wb = openpyxl.load_workbook(archivo, data_only=True)
    ws = wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _leer_filas_crudas_productos(archivo):
    """Detecta el formato por extension y devuelve la matriz cruda de celdas
    (primera fila = encabezados tal como vienen en el archivo, sin normalizar
    todavia) - paso compartido entre el preview y la importacion real."""
    ext = archivo.filename.rsplit(".", 1)[-1].lower() if "." in archivo.filename else ""
    if ext == "csv":
        return _leer_filas_csv_productos(archivo)
    elif ext in ("xlsx", "xls"):
        return _leer_filas_excel_productos(archivo)
    raise ValueError("Formato no soportado. Usa un archivo .csv o .xlsx")


def _mapeo_automatico(fila_encabezados):
    """{campo_interno: indice_columna} detectado a partir de los encabezados
    crudos del archivo, usando COLUMNAS_IMPORTACION_PRODUCTOS (con o sin
    acentos/mayusculas/espacios, ver _normalizar_header)."""
    encabezados = [_normalizar_header(h) for h in fila_encabezados]
    columnas = {}
    for idx, h in enumerate(encabezados):
        campo = COLUMNAS_IMPORTACION_PRODUCTOS.get(h)
        if campo:
            columnas[campo] = idx
    return columnas


def _leer_archivo_productos(archivo, mapeo_manual=None):
    """Parsea un archivo .csv o .xlsx de importacion masiva de productos.
    Primera fila = encabezados. `mapeo_manual`, si se pasa, es un dict
    {"indice_columna_como_texto": "campo_interno"} armado por el admin en el
    modal de preview (templates/index.html) - reemplaza por completo la
    deteccion automatica de columnas, para cuando los encabezados del archivo
    no coinciden con ningun sinonimo conocido. Retorna (filas_validas,
    errores): filas_validas es una lista de dicts normalizados (nombre,
    precio, costo, codigo, stock, stock_minimo, categoria, descripcion,
    unidad - los opcionales quedan en None si la celda venia vacia, sin
    aplicar todavia ningun default, ver database.nimbo_db.
    crear_producto_importado) y errores es una lista de {"fila": numero,
    "motivo": texto} para filas que no se pudieron importar (nombre o precio
    faltante/invalido) - no bloquean el resto del archivo."""
    filas_crudas = _leer_filas_crudas_productos(archivo)
    if not filas_crudas:
        return [], []

    if mapeo_manual:
        columnas = {}
        for idx_texto, campo in mapeo_manual.items():
            if campo:
                columnas[campo] = int(idx_texto)
    else:
        columnas = _mapeo_automatico(filas_crudas[0])

    if "nombre" not in columnas or "precio" not in columnas:
        raise ValueError("El archivo debe tener al menos las columnas Nombre y Precio.")

    def _celda(fila, campo):
        idx = columnas.get(campo)
        if idx is None or idx >= len(fila):
            return None
        valor = fila[idx]
        return None if valor is None or str(valor).strip() == "" else valor

    filas_validas, errores = [], []
    for i, fila in enumerate(filas_crudas[1:], start=2):
        nombre = _celda(fila, "nombre")
        if not nombre:
            continue  # fila vacia, se ignora sin reportar error

        nombre = str(nombre).strip()
        precio = _parsear_numero_importado(_celda(fila, "precio"))
        if precio is None or precio <= 0:
            errores.append({"fila": i, "motivo": f"'{nombre}': precio invalido o faltante"})
            continue

        codigo = _celda(fila, "codigo")
        categoria = _celda(fila, "categoria")
        descripcion = _celda(fila, "descripcion")
        unidad = _celda(fila, "unidad")

        filas_validas.append({
            "nombre": nombre,
            "precio": precio,
            "costo": _parsear_numero_importado(_celda(fila, "costo")),
            "codigo": str(codigo).strip() if codigo else None,
            "stock": _parsear_numero_importado(_celda(fila, "stock")),
            "stock_minimo": _parsear_numero_importado(_celda(fila, "stock_minimo")),
            "categoria": str(categoria).strip() if categoria else None,
            "descripcion": str(descripcion).strip() if descripcion else None,
            "unidad": str(unidad).strip() if unidad else None,
        })

    return filas_validas, errores


@app.route("/api/db/products/plantilla")
@login_required
def db_products_plantilla():
    """Template de importacion (2026-07-19): Nombre|Codigo|Precio|Costo|Stock|
    Descripcion|Unidad, con una fila de ejemplo ya llena, una fila en blanco
    lista para escribir, y una segunda hoja "Instrucciones" explicando cada
    columna - pensado tanto para el boton fijo de Inventario como para el
    boton "Descargar template" del mensaje del chat cuando la empresa
    todavia no tiene productos cargados (ver index() en app.py)."""
    import openpyxl
    from openpyxl.styles import Font
    from flask import send_file

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Productos"
    columnas = ["Nombre", "Codigo", "Precio", "Costo", "Stock", "Descripcion", "Unidad"]
    ws.append(columnas)
    for cel in ws[1]:
        cel.font = Font(bold=True)
    ws.append([
        "Licuadora Oster Clasica 3 velocidades", "7501234567890", 450, 280, 15,
        "Licuadora de vaso de vidrio, 3 velocidades, 600W", "unidad",
    ])
    ws.append(["", "", "", "", "", "", ""])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["F"].width = 40

    instrucciones = wb.create_sheet("Instrucciones")
    instrucciones.append(["Columna", "Que va aca"])
    for cel in instrucciones[1]:
        cel.font = Font(bold=True)
    filas_instrucciones = [
        ("Nombre", "Obligatoria. El nombre del producto tal como lo vas a buscar despues."),
        ("Codigo", "Opcional. El codigo de barras si ya tiene uno - si lo dejas vacio, "
                    "Nimbo le genera uno automaticamente."),
        ("Precio", "Obligatoria. El precio de venta al cliente, sin el simbolo Q."),
        ("Costo", "Opcional. Cuanto te cuesta a vos ese producto - lo usa Nimbo para "
                   "calcular margenes y sugerencias de precio."),
        ("Stock", "Opcional. Cuantas unidades tenes ahorita mismo. Si lo dejas vacio, "
                   "entra con 0."),
        ("Descripcion", "Opcional. Una descripcion mas larga del producto, para mostrarla "
                         "en el chat o en la pagina publica del producto."),
        ("Unidad", "Opcional. Como se vende (unidad, caja, docena, par...). Si lo dejas "
                    "vacio, queda como 'unidad'."),
    ]
    for fila in filas_instrucciones:
        instrucciones.append(fila)
    instrucciones.column_dimensions["A"].width = 14
    instrucciones.column_dimensions["B"].width = 80

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name="template_inventario_nimbo.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.route("/api/db/products/import/preview", methods=["POST"])
@login_required
def db_products_import_preview():
    """Paso previo a importar de verdad: lee el archivo, sugiere el mapeo de
    columnas automatico (mismo criterio que la importacion real) y devuelve
    los encabezados crudos + las primeras 5 filas de datos para que el modal
    de Inventario (o el que abre el boton 'Subir mi Excel' del chat) las
    muestre antes de confirmar - no importa nada todavia."""
    if "archivo" not in request.files:
        return jsonify({"error": "No se envio ningun archivo"}), 400
    archivo = request.files["archivo"]
    if archivo.filename == "":
        return jsonify({"error": "Archivo vacio"}), 400

    try:
        filas_crudas = _leer_filas_crudas_productos(archivo)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"No se pudo leer el archivo: {e}"}), 400

    if not filas_crudas:
        return jsonify({"error": "El archivo está vacío."}), 400

    headers_crudos = ["" if h is None else str(h) for h in filas_crudas[0]]
    mapeo_sugerido = {
        str(idx): campo for campo, idx in _mapeo_automatico(filas_crudas[0]).items()
    }
    filas_datos = filas_crudas[1:]
    filas_preview = [
        ["" if celda is None else celda for celda in fila]
        for fila in filas_datos[:5]
    ]

    return jsonify({
        "headers": headers_crudos,
        "mapeo_sugerido": mapeo_sugerido,
        "campos_disponibles": CAMPOS_IMPORTACION_LABELS,
        "filas_preview": filas_preview,
        "total_filas": len(filas_datos),
    })


@app.route("/api/db/products/import", methods=["POST"])
@login_required
def db_products_import():
    from database.nimbo_db import crear_producto_importado, buscar_producto_duplicado
    if "archivo" not in request.files:
        return jsonify({"error": "No se envio ningun archivo"}), 400
    archivo = request.files["archivo"]
    if archivo.filename == "":
        return jsonify({"error": "Archivo vacio"}), 400

    # Mapeo manual opcional (JSON, campo de formulario "mapeo") armado en el
    # paso de preview del modal - si no viene, _leer_archivo_productos cae a
    # la deteccion automatica de siempre (mismo comportamiento que antes de
    # este feature).
    mapeo_manual = None
    mapeo_raw = request.form.get("mapeo")
    if mapeo_raw:
        try:
            mapeo_manual = json.loads(mapeo_raw)
        except (TypeError, ValueError):
            mapeo_manual = None

    try:
        filas, errores = _leer_archivo_productos(archivo, mapeo_manual)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"No se pudo leer el archivo: {e}"}), 400

    creados = 0
    duplicados = []
    for fila in filas:
        existente = buscar_producto_duplicado(fila["nombre"], fila.get("codigo"))
        if existente:
            duplicados.append({"fila": fila, "producto_existente": existente})
            continue
        crear_producto_importado(fila)
        creados += 1

    return jsonify({"creados": creados, "duplicados": duplicados, "errores": errores})


@app.route("/api/db/products/import/resolver", methods=["POST"])
@login_required
def db_products_import_resolver():
    from database.nimbo_db import crear_producto_importado, actualizar_producto_importado
    data = request.json or {}
    resoluciones = data.get("resoluciones", [])

    creados, actualizados, saltados = 0, 0, 0
    for r in resoluciones:
        accion = r.get("accion")
        fila = r.get("fila")
        if not fila:
            continue
        if accion == "actualizar":
            producto_id = (r.get("producto_existente") or {}).get("Id")
            if not producto_id:
                continue
            actualizar_producto_importado(producto_id, fila)
            actualizados += 1
        elif accion == "crear":
            crear_producto_importado(fila)
            creados += 1
        else:
            saltados += 1

    return jsonify({"creados": creados, "actualizados": actualizados, "saltados": saltados})


@app.route("/api/db/products/<int:pid>/etiqueta")
@login_required
def db_product_etiqueta(pid):
    from database.nimbo_db import get_product_by_id, asignar_codigo_barras
    from services.barcode_service import generar_etiqueta_producto_pdf
    producto = get_product_by_id(pid)
    if not producto:
        return jsonify({"error": "Producto no encontrado"}), 404
    nombre = producto["Name"].strip()
    codigo = producto.get("Code") or asignar_codigo_barras(pid, nombre)
    ruta = generar_etiqueta_producto_pdf(nombre, producto["Price"], codigo)
    return jsonify({"ok": True, "ruta": f"/static/{ruta}"})


@app.route("/api/db/stock/add", methods=["POST"])
@login_required
def db_add_stock():
    from database.nimbo_db import add_stock
    data = request.json
    ok = add_stock(data["product_id"], data["quantity"])
    return jsonify({"ok": ok})

@app.route("/api/db/stock/remove", methods=["POST"])
@login_required
def db_remove_stock():
    from database.nimbo_db import remove_stock
    data = request.json
    ok = remove_stock(data["product_id"], data["quantity"])
    if not ok:
        return jsonify({"error": "Stock insuficiente"}), 400
    return jsonify({"ok": ok})

@app.route("/api/db/stock/<int:pid>", methods=["GET"])
@login_required
def db_get_stock(pid):
    from database.nimbo_db import get_stock
    qty = get_stock(pid)
    return jsonify({"stock": qty})


@app.route("/api/db/kardex")
@login_required
def db_kardex():
    from database.nimbo_db import get_kardex
    product_id = request.args.get("product_id")
    limit = int(request.args.get("limit", 50))
    return jsonify(get_kardex(product_id=product_id, limit=limit))

@app.route("/api/db/alertas")
@login_required
def db_alertas():
    from database.nimbo_db import get_alertas_stock
    return jsonify(get_alertas_stock())

@app.route("/api/inventario/alertas-predictivas")
@login_required
def db_alertas_predictivas():
    from database.nimbo_db import get_alertas_predictivas
    return jsonify(get_alertas_predictivas())

@app.route("/api/db/stock-minimo", methods=["POST"])
@login_required
def db_set_stock_minimo():
    from database.nimbo_db import set_stock_minimo
    data = request.json
    set_stock_minimo(data["product_id"], data["product_name"], data["minimo"])
    return jsonify({"ok": True})


@app.route("/api/download/<path:filename>")
@login_required
def download_file(filename):
    from flask import send_file
    filepath = os.path.join(BASE_DIR, filename)
    if os.path.exists(filepath):
        return send_file(filepath, as_attachment=True)
    return jsonify({"error": "Archivo no encontrado"}), 404


@app.route("/api/db/alertas-precios")
@admin_required
def db_alertas_precios():
    from database.nimbo_db import get_alertas_precios
    return jsonify(get_alertas_precios())


UPLOAD_FOLDER = os.path.join(BASE_DIR, 'static', 'fotos')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

@app.route("/api/productos/foto", methods=["POST"])
@login_required
def subir_foto_producto():
    from flask import request
    from database.nimbo_db import nimbo_connect
    
    if 'foto' not in request.files:
        return jsonify({"error": "No se envio foto"}), 400
    
    producto_nombre = request.form.get("producto", "")
    producto_id = request.form.get("producto_id", 0)
    foto = request.files['foto']
    
    if foto.filename == '':
        return jsonify({"error": "Archivo vacio"}), 400
    
    ext = foto.filename.rsplit('.', 1)[-1].lower()
    if ext not in ['jpg', 'jpeg', 'png', 'webp']:
        return jsonify({"error": "Formato no permitido"}), 400
    
    import uuid
    filename = f"{uuid.uuid4()}.{ext}"
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    foto.save(filepath)
    
    conn = nimbo_connect()
    cur = conn.cursor()
    cur.execute("INSERT INTO ProductoFoto (ProductoId, ProductoNombre, RutaFoto) VALUES (?,?,?)",
        (producto_id, producto_nombre, f"/static/fotos/{filename}"))
    conn.commit()
    conn.close()
    
    return jsonify({"ok": True, "ruta": f"/static/fotos/{filename}"})

@app.route("/api/productos/fotos/<int:producto_id>")
@login_required
def get_fotos_producto(producto_id):
    from database.nimbo_db import nimbo_connect
    conn = nimbo_connect()
    cur = conn.cursor()
    cur.execute("SELECT * FROM ProductoFoto WHERE ProductoId=? ORDER BY FechaSubida DESC", (producto_id,))
    fotos = [dict(r) for r in cur.fetchall()]
    conn.close()
    return jsonify(fotos)


UPLOAD_FOLDER_SERVICIO = os.path.join(BASE_DIR, 'static', 'fotos_servicio')

@app.route("/api/servicio/fotos", methods=["POST"])
@admin_o_tecnico_required
def subir_foto_servicio():
    from database.nimbo_db import get_orden, agregar_foto_orden

    if 'foto' not in request.files:
        return jsonify({"error": "No se envio foto"}), 400

    numero_orden = request.form.get("numero_orden", "").strip().upper()
    if not numero_orden.startswith("OS-"):
        numero_orden = "OS-" + numero_orden.zfill(4)
    descripcion = request.form.get("descripcion", "")
    foto = request.files['foto']

    if foto.filename == '':
        return jsonify({"error": "Archivo vacio"}), 400

    orden = get_orden(numero_orden)
    if not orden:
        return jsonify({"error": f"No encontre la orden {numero_orden}"}), 404

    ext = foto.filename.rsplit('.', 1)[-1].lower()
    if ext not in ['jpg', 'jpeg', 'png', 'webp']:
        return jsonify({"error": "Formato no permitido"}), 400

    carpeta_orden = os.path.join(UPLOAD_FOLDER_SERVICIO, numero_orden)
    os.makedirs(carpeta_orden, exist_ok=True)

    import uuid
    filename = f"{uuid.uuid4()}.{ext}"
    foto.save(os.path.join(carpeta_orden, filename))

    ruta = f"/static/fotos_servicio/{numero_orden}/{filename}"
    agregar_foto_orden(orden["Id"], numero_orden, ruta, descripcion)

    return jsonify({"ok": True, "ruta": ruta})


@app.route("/api/servicio/fotos/<numero_orden>")
@admin_o_tecnico_required
def get_fotos_servicio(numero_orden):
    from database.nimbo_db import get_fotos_orden
    numero_orden = numero_orden.strip().upper()
    if not numero_orden.startswith("OS-"):
        numero_orden = "OS-" + numero_orden.zfill(4)
    return jsonify(get_fotos_orden(numero_orden))


@app.route("/api/servicio/ordenes")
@admin_o_tecnico_required
def get_ordenes_servicio():
    from database.nimbo_db import get_ordenes
    estado = request.args.get("estado")
    return jsonify(get_ordenes(estado=estado, limit=100))


@app.route("/api/servicio/pipeline")
@admin_o_tecnico_required
def get_pipeline_servicio_route():
    from database.nimbo_db import get_pipeline_servicio
    return jsonify(get_pipeline_servicio())


@app.route("/producto/<codigo>")
def ver_producto_publico(codigo):
    """Vista publica (sin login) a la que apunta el QR de las etiquetas de producto -
    ver services/barcode_service.py. Muestra solo foto/nombre/precio/descripcion,
    nunca stock ni costo."""
    from database.nimbo_db import get_product_by_code, get_foto_producto, get_descripcion_producto

    producto = get_product_by_code(codigo)
    if not producto:
        return render_template("producto_publico.html", encontrado=False), 404

    info = get_descripcion_producto(producto["Id"])
    descripcion = (info["Descripcion"] if info and info.get("Descripcion") else "") or (producto.get("Description") or "")
    foto = get_foto_producto(producto["Id"])

    return render_template("producto_publico.html",
        encontrado=True,
        nombre=producto["Name"].strip(),
        precio=producto["Price"],
        descripcion=descripcion.strip(),
        foto=foto)


@app.route("/tiktok-verification.html")
def tiktok_verify():
    return "tiktok-verification"


@app.route("/tiktokpxma0DpKg90IGKsTJYXp0AInP21T3LWN.txt")
@app.route("/tiktokpxma0DpKg90IGKsTJYXp0AInP21T3LWN.txt/")
def tiktok_verify_file():
    return "tiktok-developers-site-verification=pxma0DpKg90IGKsTJYXp0AInP21T3LWN"


@app.route("/api/limpieza-anual", methods=["POST"])
@admin_required
def limpieza_anual():
    from database.nimbo_db import archivar_y_limpiar_anio
    data = request.json
    anio = data.get("anio", 2026)
    try:
        archivo = archivar_y_limpiar_anio(anio)
        return jsonify({"ok": True, "archivo": archivo})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/recordatorios")
@admin_required
def get_recordatorios():
    from database.nimbo_db import get_recordatorios
    return jsonify(get_recordatorios())


@app.route("/api/cotizaciones/resumen")
@admin_required
def get_resumen_cotizaciones():
    from database.nimbo_db import get_resumen_cotizaciones
    return jsonify(get_resumen_cotizaciones())


@app.route("/api/cuentascobrar/resumen")
@admin_required
def get_resumen_cuentas_cobrar_route():
    from database.nimbo_db import get_resumen_cuentas_cobrar
    return jsonify(get_resumen_cuentas_cobrar())


@app.route("/api/cuentaspagar/resumen")
@admin_required
def get_resumen_cuentas_pagar_route():
    from database.nimbo_db import get_resumen_cuentas_pagar
    return jsonify(get_resumen_cuentas_pagar())


@app.route("/api/contabilidad/flujo-caja-proyectado")
@admin_required
def get_flujo_caja_proyectado_route():
    from database.nimbo_db import get_flujo_caja_proyectado
    return jsonify(get_flujo_caja_proyectado())


@app.route("/api/reportes/meta-semanal", methods=["GET", "POST"])
@admin_required
def meta_semanal_route():
    from database.nimbo_db import get_progreso_meta_semanal, set_meta_semanal
    if request.method == "POST":
        data = request.json or {}
        try:
            monto = float(data.get("monto"))
            if monto <= 0:
                raise ValueError()
        except (TypeError, ValueError):
            return jsonify({"error": "Monto invalido"}), 400
        return jsonify(set_meta_semanal(monto))
    return jsonify(get_progreso_meta_semanal() or {})


# ── CAJA REGISTRADORA ──────────────────────────────────────

def _requiere_turno_abierto():
    from database.nimbo_db import get_turno_abierto
    turno = get_turno_abierto()
    if not turno:
        return None, (jsonify({"error": "No hay ningun turno de caja abierto."}), 400)
    return turno, None


@app.route("/api/caja/estado")
@login_required
def caja_estado():
    from database.nimbo_db import (
        get_turno_abierto, get_denominaciones_actuales, sencillo_bajo, DENOMINACIONES,
        get_ultimo_cierre_denominaciones, get_fondo_sugerido, get_caja_fuerte_acumulada
    )
    turno = get_turno_abierto()
    if not turno:
        return jsonify({
            "turno": None,
            "denominaciones": DENOMINACIONES,
            "fondo_sugerido": get_ultimo_cierre_denominaciones(),
            "fondo_sugerido_total": get_fondo_sugerido(),
            "caja_fuerte_acumulada": get_caja_fuerte_acumulada()
        })
    bajo, pequenas = sencillo_bajo(turno["Id"])
    return jsonify({
        "turno": turno,
        "denominaciones_actuales": get_denominaciones_actuales(turno["Id"]),
        "denominaciones": DENOMINACIONES,
        "sencillo_bajo": bajo,
        "denominaciones_pequenas": pequenas,
        "caja_fuerte_acumulada": get_caja_fuerte_acumulada()
    })


@app.route("/api/caja/abrir", methods=["POST"])
@login_required
def caja_abrir():
    """El fondo inicial se toma automaticamente de Configuracion.caja_fondo_sugerido
    (lo que quedo contado al cerrar el turno anterior, o el fondo configurado por un
    admin si todavia no se ha cerrado ningun turno - ver get_fondo_sugerido). El
    empleado solo confirma la apertura, no puede enviar un monto distinto: cambiar el
    fondo es exclusivo de Configuracion del negocio (solo admin)."""
    from database.nimbo_db import get_turno_abierto, abrir_turno_simple, get_fondo_sugerido
    if get_turno_abierto():
        return jsonify({"error": "Ya hay un turno de caja abierto."}), 400
    fondo_total = get_fondo_sugerido()
    if fondo_total is None:
        fondo_total = 0
    turno_id, total = abrir_turno_simple(int(current_user.id), current_user.nombre, fondo_total)
    return jsonify({"ok": True, "turno_id": turno_id, "efectivo_esperado": total})


@app.route("/api/caja/cerrar", methods=["POST"])
@login_required
def caja_cerrar():
    from database.nimbo_db import cerrar_turno
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    data = request.json or {}
    try:
        contado_total = float(data.get("contado_total"))
        if contado_total < 0:
            raise ValueError()
    except (TypeError, ValueError):
        return jsonify({"error": "Escribe el monto total contado en caja."}), 400
    try:
        retiro_caja_fuerte = float(data.get("retiro_caja_fuerte", 0) or 0)
        if retiro_caja_fuerte < 0:
            raise ValueError()
    except (TypeError, ValueError):
        return jsonify({"error": "Monto de retiro invalido."}), 400
    if retiro_caja_fuerte > contado_total:
        return jsonify({"error": "El retiro no puede ser mayor a lo contado en caja."}), 400
    conteo_fisico = data.get("conteo_fisico") or None
    resultado = cerrar_turno(turno["Id"], contado_total, retiro_caja_fuerte, conteo_fisico)
    return jsonify(resultado)


@app.route("/api/caja/carrito")
@login_required
def caja_carrito():
    from database.nimbo_db import get_carrito
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    items = get_carrito(turno["Id"])
    total = round(sum(i["Total"] for i in items), 2)
    return jsonify({"items": items, "total": total})


@app.route("/api/caja/carrito/agregar", methods=["POST"])
@login_required
def caja_carrito_agregar():
    from database.nimbo_db import get_product_by_id, get_product_by_code, get_products, agregar_a_carrito
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    data = request.json or {}
    cantidad = data.get("cantidad", 1)
    try:
        cantidad = float(cantidad)
    except (TypeError, ValueError):
        return jsonify({"error": "Cantidad invalida"}), 400
    if cantidad <= 0:
        return jsonify({"error": "Cantidad invalida"}), 400

    producto = None
    if data.get("producto_id"):
        producto = get_product_by_id(data["producto_id"])
    elif data.get("query"):
        query = data["query"].strip()
        producto = get_product_by_code(query)
        if not producto:
            productos = get_products(search=query)
            if len(productos) > 1:
                return jsonify({"multiples": True, "productos": productos[:8]})
            producto = productos[0] if productos else None

    if not producto:
        return jsonify({"error": "Producto no encontrado."}), 404
    if producto["Stock"] < cantidad:
        return jsonify({"error": f"Stock insuficiente de {producto['Name']} ({int(producto['Stock'])} disponibles)."}), 400

    items = agregar_a_carrito(turno["Id"], producto["Id"], producto["Name"], cantidad, producto["Price"])
    total = round(sum(i["Total"] for i in items), 2)
    return jsonify({"items": items, "total": total})


@app.route("/api/caja/carrito/quitar", methods=["POST"])
@login_required
def caja_carrito_quitar():
    from database.nimbo_db import quitar_de_carrito
    data = request.json or {}
    item_id = data.get("item_id")
    if not item_id:
        return jsonify({"error": "item_id requerido"}), 400
    items = quitar_de_carrito(item_id)
    total = round(sum(i["Total"] for i in items), 2)
    return jsonify({"items": items, "total": total})


@app.route("/api/caja/carrito/vaciar", methods=["POST"])
@login_required
def caja_carrito_vaciar():
    from database.nimbo_db import vaciar_carrito
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    vaciar_carrito(turno["Id"])
    return jsonify({"ok": True, "items": [], "total": 0})


@app.route("/api/caja/verificar-cobrador", methods=["POST"])
@login_required
def caja_verificar_cobrador():
    """Popup de identificacion que aparece antes de cada cobro en Caja (no solo al
    abrir turno): valida usuario/contrasena y que tenga permiso de caja, para
    registrar quien cobro esa venta puntual (Venta.CobradorNombre)."""
    data = request.json or {}
    usuario = verify_login(data.get("username", ""), data.get("password", ""))
    if not usuario:
        return jsonify({"error": "Usuario o contraseña incorrectos."}), 401
    if not usuario.tiene_permiso("caja"):
        return jsonify({"error": "No tienes permiso para cobrar."}), 403
    return jsonify({"ok": True, "nombre": usuario.nombre})


@app.route("/api/caja/cobrar", methods=["POST"])
@login_required
def caja_cobrar():
    from database.nimbo_db import cobrar_carrito
    from services.ticket import generar_texto_ticket
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    data = request.json or {}
    cobrador_nombre = data.get("cobrador_nombre") or current_user.nombre
    resultado = cobrar_carrito(
        turno["Id"],
        data.get("metodo_pago"),
        cliente_id=data.get("cliente_id"),
        monto_efectivo=data.get("monto_efectivo"),
        monto_tarjeta=data.get("monto_tarjeta"),
        denominaciones_recibidas=data.get("denominaciones_recibidas"),
        cobrador_nombre=cobrador_nombre
    )
    if resultado.get("error"):
        return jsonify(resultado), 400

    ticket_items = [{
        "nombre": i["ProductoNombre"], "cantidad": i["Cantidad"],
        "precio_unit": i["PrecioUnit"], "subtotal": i["Total"]
    } for i in resultado["items"]]
    resultado["ticket"] = generar_texto_ticket(
        resultado["venta_id"], resultado["fecha"], ticket_items, resultado["total"],
        resultado["metodo_pago"], cliente_nombre=data.get("cliente_nombre", ""),
        efectivo_recibido=resultado["monto_efectivo"] if resultado["metodo_pago"] == "efectivo" else None,
        cambio=resultado["cambio"], cajero=cobrador_nombre,
        subtotal=resultado["subtotal"], descuento=resultado["descuento"],
        garantias=resultado["garantias"]
    )
    return jsonify(resultado)


@app.route("/api/caja/retirar", methods=["POST"])
@login_required
def caja_retirar():
    if not current_user.is_admin:
        return jsonify({"error": "Solo un administrador puede retirar efectivo de caja."}), 403
    from database.nimbo_db import retirar_caja
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    data = request.json or {}
    monto = data.get("monto", 0)
    try:
        monto = float(monto)
    except (TypeError, ValueError):
        return jsonify({"error": "Monto invalido"}), 400
    resultado = retirar_caja(turno["Id"], monto, data.get("descripcion", ""), data.get("denominaciones"))
    if resultado.get("error"):
        return jsonify(resultado), 400
    return jsonify(resultado)


@app.route("/api/caja/anular", methods=["POST"])
@login_required
def caja_anular():
    if not current_user.is_admin:
        return jsonify({"error": "Solo un administrador puede anular un turno de caja."}), 403
    from database.nimbo_db import anular_turno
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    data = request.json or {}
    motivo = (data.get("motivo") or "").strip()
    if not motivo:
        return jsonify({"error": "Escribe el motivo de la anulacion."}), 400
    resultado = anular_turno(turno["Id"], motivo, current_user.nombre)
    if resultado.get("error"):
        return jsonify(resultado), 400
    return jsonify(resultado)


@app.route("/api/caja/movimientos")
@login_required
def caja_movimientos():
    from database.nimbo_db import get_movimientos_caja
    turno, error = _requiere_turno_abierto()
    if error:
        return error
    return jsonify(get_movimientos_caja(turno["Id"]))


@app.route("/api/caja/resumen-dia")
@login_required
def caja_resumen_dia():
    from database.nimbo_db import get_resumen_caja_dia
    return jsonify(get_resumen_caja_dia())


@app.route("/api/caja/imprimir-ticket", methods=["POST"])
@login_required
def caja_imprimir_ticket():
    from database.nimbo_db import nimbo_connect
    from services.ticket import generar_texto_ticket
    data = request.json or {}
    venta_id = data.get("venta_id")
    if not venta_id:
        return jsonify({"error": "venta_id requerido"}), 400
    conn = nimbo_connect()
    cur = conn.cursor()
    cur.execute("""SELECT v.*, c.Nombre as ClienteNombre FROM Venta v
        LEFT JOIN Cliente c ON v.ClienteId=c.Id WHERE v.Id=?""", (venta_id,))
    row = cur.fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "Venta no encontrada"}), 404
    venta = dict(row)
    # Reimpresion: no se conserva el detalle por producto de ventas ya cerradas,
    # se reconstruye un ticket simplificado a partir de la descripcion guardada.
    items = [{"nombre": venta["Descripcion"], "cantidad": 1, "precio_unit": venta["Total"], "subtotal": venta["Total"]}]
    ticket = generar_texto_ticket(
        venta_id, venta["Fecha"], items, venta["Total"], "(reimpresion)",
        cliente_nombre=venta.get("ClienteNombre") or ""
    )
    return jsonify({"ticket": ticket})


@app.route("/api/chat/historial")
@admin_required
def get_historial_guardado():
    from database.nimbo_db import get_historial_chat
    return jsonify(get_historial_chat())

@app.route("/api/chat/limpiar-historial", methods=["POST"])
@admin_required
def limpiar_historial_guardado():
    from database.nimbo_db import limpiar_historial_chat
    limpiar_historial_chat()
    return jsonify({"ok": True})


@app.route("/api/chat/archivar", methods=["POST"])
@admin_required
def archivar_chat():
    from database.nimbo_db import archivar_chat_trimestral
    archivo = archivar_chat_trimestral()
    if archivo:
        return jsonify({"ok": True, "archivo": archivo.split("/")[-1]})
    return jsonify({"ok": True, "mensaje": "No hay mensajes antiguos para archivar"})


@app.route("/api/contabilidad/equilibrio")
@admin_required
def get_equilibrio():
    from database.nimbo_db import get_gastos, get_pagos_fijos, get_margen_productos
    gastos = get_gastos()
    total_gastos = sum(float(g["Monto"]) for g in gastos)
    productos = get_margen_productos()
    if not productos:
        return jsonify({"error": "Sin datos"})
    precios = [float(p["Price"]) for p in productos if float(p["Price"]) > 0]
    costos = [float(p["Cost"]) for p in productos if float(p["Cost"]) > 0]
    if not precios or not costos:
        return jsonify({"error": "Sin precios"})
    precio_prom = sum(precios) / len(precios)
    costo_prom = sum(costos) / len(costos)
    contribucion = (precio_prom - costo_prom) - (precio_prom * 0.06)
    if contribucion <= 0:
        return jsonify({"error": "Margen negativo"})
    unidades = total_gastos / contribucion if total_gastos > 0 else 0
    ventas_necesarias = unidades * precio_prom
    from database.nimbo_db import nimbo_connect
    from datetime import datetime
    conn = nimbo_connect()
    cur = conn.cursor()
    cur.execute("SELECT IFNULL(SUM(Total),0) as total FROM Venta WHERE Mes=? AND Anio=?",
        (datetime.now().month, datetime.now().year))
    ventas_actuales = float(cur.fetchone()["total"])
    conn.close()
    return jsonify({
        "gastos": round(total_gastos, 2),
        "ventas_necesarias": round(ventas_necesarias, 2),
        "ventas_actuales": round(ventas_actuales, 2),
        "superado": ventas_actuales >= ventas_necesarias
    })


@app.route("/api/graficas/ventas")
@admin_required
def graficas_ventas():
    from database.nimbo_db import get_comparacion_meses
    meses = get_comparacion_meses(6)
    return jsonify({
        "labels": [m["nombre"] for m in meses],
        "ventas": [m["total"] for m in meses],
        "gastos": [m["gastos"] for m in meses],
        "utilidad": [m["utilidad"] for m in meses]
    })


@app.route("/api/backup", methods=["POST"])
@admin_required
def hacer_backup_route():
    from database.nimbo_db import hacer_backup
    try:
        fecha = hacer_backup()
        return jsonify({"ok": True, "fecha": fecha})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    print("=" * 50)
    print("  Nimbo AI Dashboard")
    print("  http://localhost:5000")
    print("=" * 50)
    app.run(debug=True, port=5000)
